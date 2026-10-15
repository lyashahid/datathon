"""Read-only inventory and grain checks for local analytical inputs.

Run with uv run --directory backend python ../scripts/analysis/inventory-review.py.
Requires the optional original archive files listed by --check-inputs.
Writes only derived JSON/CSV artifacts under artifacts/analysis, never application inputs.
"""

from __future__ import annotations

import argparse
import copy
import csv
import hashlib
import json
import math
import sys
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))
from app.data.normalize.names import normalize_atoll_code, normalize_name
from app.data.population import ATOLLS, attach_population, population_summary


def read(path):
    return json.loads((ROOT / path).read_text(encoding="utf-8"))


def receipt(path):
    p = ROOT / path
    content = p.read_bytes()
    return {
        "path": p.relative_to(ROOT).as_posix(),
        "sha256": hashlib.sha256(content).hexdigest(),
        "bytes": len(content),
    }


def counts(rows, key):
    return dict(Counter(str(row.get(key)) for row in rows))


def unique_check(rows, key="id"):
    return {
        "rows": len(rows),
        "distinct_keys": len({r.get(key) for r in rows}),
        "missing_keys": sum(r.get(key) is None for r in rows),
    }


def extract_island_indicator(path, islands):
    wb = openpyxl.load_workbook(ROOT / path, read_only=True, data_only=True)
    sheet = wb.active
    rows = []
    for row_no, values in enumerate(sheet.iter_rows(values_only=True), 1):
        code = normalize_atoll_code(values[0])
        if code in ATOLLS and values[1] and isinstance(values[2], (int, float)):
            rows.append(
                {
                    "source_row": row_no,
                    "atoll_code": values[0],
                    "locality": values[1],
                    "source": "island-indicator",
                    "values": list(values[2:]),
                }
            )
    wb.close()
    attach_population(rows, islands)
    return rows


def finite_number(value):
    return (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and math.isfinite(value)
    )


ARCHIVE_INPUTS = (
    "data/raw/census2022-population.xlsx",
    "data/raw/census2022-employment.xlsx",
    "data/raw/island-projections2022-2030.xlsx",
    "data/raw/population-definitions.xlsx",
    "data/raw/employment-definitions.xlsx",
    "data/public/archive/policy-evidence/igmh-utilisation.xlsx",
    "data/public/archive/policy-evidence/school-statistics-2023.pdf",
    "data/public/archive/policy-evidence/constituency-population.xlsx",
    "data/public/archive/policy-evidence/constituency-employment.xlsx",
    "data/public/archive/policy-evidence/constituency-education.xlsx",
    "backend/data/cache/raw/mbs/P3.xlsx",
    "backend/data/cache/raw/mbs/P4.xlsx",
    "backend/data/cache/raw/mbs/P5.xlsx",
    "backend/data/cache/raw/mbs/P6.xlsx",
)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "artifacts" / "analysis" / datetime.now(UTC).date().isoformat(),
        help="Derived output directory; defaults to artifacts/analysis/<UTC run date>.",
    )
    parser.add_argument(
        "--check-inputs",
        action="store_true",
        help="List required original archive files and exit without reading or writing outputs.",
    )
    args = parser.parse_args()
    missing = [name for name in ARCHIVE_INPUTS if not (ROOT / name).is_file()]
    if args.check_inputs:
        print(
            json.dumps({"archiveInputs": ARCHIVE_INPUTS, "missing": missing}, indent=2)
        )
        raise SystemExit(1 if missing else 0)
    if missing:
        parser.exit(
            2,
            "Optional historical inventory audit requires original source archives. "
            "The application and forecast audit use bundled data and do not need these files.\n"
            "Missing archive files:\n"
            + "\n".join(f"  {name}" for name in missing)
            + "\nSee docs/data-analysis-opportunities.md (Evidence and reproducibility).\n",
        )
    OUT = args.output.resolve()
    islands = read("backend/data/cache/normalized/islands.geojson")["features"]
    island_ids = {x["properties"]["id"] for x in islands}
    island_properties = {x["properties"]["id"]: x["properties"] for x in islands}
    population_stored = read("backend/data/cache/normalized/population.json")
    population = copy.deepcopy(population_stored)
    attach_population(population, islands)
    population_by_island = defaultdict(list)
    for record in population:
        if record.get("island_id"):
            population_by_island[record["island_id"]].append(record)
    island_census = {
        key: population_summary(rows) for key, rows in population_by_island.items()
    }
    facilities = read("backend/data/cache/normalized/facilities.json")
    transport = read("backend/data/cache/normalized/transport.json")
    baseline = read("data/public/planning-baseline.json")
    baseline_node_ids = {r["id"] for r in baseline["network"]["nodes"]}
    policy = read("data/public/policy-evidence/bundle.json")
    context = read("data/public/context/agency-context.json")
    operations = read("data/public/operations/bundle.json")
    prediction = read("data/public/prediction-source.json")
    calendars = read("data/public/policy-evidence/service-calendars.json")
    pop_indicator = extract_island_indicator(
        "data/raw/census2022-population.xlsx", islands
    )
    employment = extract_island_indicator(
        "data/raw/census2022-employment.xlsx", islands
    )
    p4 = openpyxl.load_workbook(
        ROOT / "backend/data/cache/raw/mbs/P4.xlsx", read_only=True, data_only=True
    )
    national_census_total = p4.active.cell(7, 3).value
    assert p4.active.cell(7, 2).value == "Republic" and national_census_total == 515132
    p4.close()

    population_numeric_gaps = [
        {
            "row": x["source_row"],
            "atoll": x["atoll_code"],
            "locality": x["locality"],
            "column_1_based": i + 3,
            "value": v,
        }
        for x in pop_indicator
        for i, v in enumerate(x["values"])
        if not finite_number(v)
    ]
    employment_numeric_gaps = [
        {
            "row": x["source_row"],
            "atoll": x["atoll_code"],
            "locality": x["locality"],
            "column_1_based": i + 3,
            "value": v,
        }
        for x in employment
        for i, v in enumerate(x["values"])
        if not finite_number(v)
    ]
    pop_reconciliation_failures = []
    for row in pop_indicator:
        v = row["values"]
        checks = {
            "resident_sex": v[0] == v[1] + v[2],
            "nationality": v[0] == v[3] + v[6],
            "maldivian_age_partition": v[3] == v[21] + v[27] + v[36],
        }
        for name, valid in checks.items():
            if not valid:
                pop_reconciliation_failures.append(
                    {
                        "row": row["source_row"],
                        "check": name,
                        "locality": row["locality"],
                    }
                )

    OUT.mkdir(parents=True, exist_ok=True)
    with (OUT / "island-indicator-review.csv").open(
        "w", newline="", encoding="utf-8"
    ) as handle:
        fields = [
            "source_row",
            "island_id",
            "locality",
            "atoll_code",
            "resident_2022",
            "maldivian_2022",
            "foreign_2022",
            "maldivian_2014",
            "maldivian_growth_2014_2022_percent",
            "maldivian_child_0_14",
            "maldivian_working_age_15_64",
            "maldivian_65_plus",
            "maldivian_dependency_ratio",
            "resident_density_per_km2",
            "area_ha",
        ]
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for x in pop_indicator:
            v = x["values"]
            writer.writerow(
                dict(
                    zip(
                        fields,
                        [
                            x["source_row"],
                            x["island_id"],
                            x["locality"],
                            x["atoll_code"],
                            v[0],
                            v[3],
                            v[6],
                            v[9],
                            v[12],
                            v[21],
                            v[27],
                            v[36],
                            v[15],
                            v[48],
                            v[46],
                        ],
                    )
                )
            )
    with (OUT / "island-employment-review.csv").open(
        "w", newline="", encoding="utf-8"
    ) as handle:
        fields = [
            "source_row",
            "island_id",
            "locality",
            "atoll_code",
            "match_status",
            "resident_lfp_total_percent",
            "resident_lfp_women_percent",
            "resident_lfp_men_percent",
            "resident_unemployment_total_percent",
            "maldivian_lfp_total_percent",
            "maldivian_lfp_women_percent",
            "maldivian_lfp_men_percent",
            "maldivian_neet_15_24_percent",
            "maldivian_neet_18_35_percent",
        ]
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for x in employment:
            v = x["values"]
            writer.writerow(
                dict(
                    zip(
                        fields,
                        [
                            x["source_row"],
                            x["island_id"],
                            x["locality"],
                            x["atoll_code"],
                            x["match_status"],
                            v[0],
                            v[1],
                            v[2],
                            v[6],
                            v[27],
                            v[28],
                            v[29],
                            v[48],
                            v[51],
                        ],
                    )
                )
            )

    layer_profiles = []
    for layer in policy["layers"]:
        features = layer["features"]["features"]
        properties = [x["properties"] for x in features]
        keys = set().union(*(r.keys() for r in properties))
        categorical = {
            k: counts(properties, k)
            for k in keys
            if k
            in {
                "landuse",
                "phase",
                "HEALTH",
                "MATURITY",
                "ramp_available",
                "step_free_entrance",
                "continuous_movement",
                "pathway_width",
                "door_accessible",
                "surface_condition",
                "category",
                "ownership",
            }
        }
        dates = {}
        for k in {
            "created_date",
            "last_edited_date",
            "CREATED_DATE",
            "LAST_EDITED_DATE",
        } & keys:
            values = [r[k] for r in properties if finite_number(r.get(k))]
            dates[k] = {
                "populated": len(values),
                "min": datetime.fromtimestamp(min(values) / 1000, UTC).isoformat()
                if values
                else None,
                "max": datetime.fromtimestamp(max(values) / 1000, UTC).isoformat()
                if values
                else None,
            }
        layer_profiles.append(
            {
                "id": layer["id"],
                "reference_period": layer.get("referencePeriod"),
                "record_count": len(features),
                "original_count": layer.get("sourceRecordCount"),
                "unlocated": layer.get("unlocatedRecords"),
                "key_check": unique_check(properties),
                "categorical_counts": categorical,
                "date_coverage": dates,
                "null_counts": {
                    k: sum(r.get(k) in (None, "") for r in properties)
                    for k in sorted(keys)
                },
            }
        )

    igmh_path = "data/public/archive/policy-evidence/igmh-utilisation.xlsx"
    wb = openpyxl.load_workbook(ROOT / igmh_path, read_only=True, data_only=True)
    sheet = wb.active
    igmh = []
    for column in range(2, 23):
        year = sheet.cell(4, column).value
        for source_row in range(6, 12):
            value = sheet.cell(source_row, column).value
            if value is not None:
                igmh.append(
                    {
                        "year": year,
                        "metric": sheet.cell(source_row, 1).value,
                        "value": value if finite_number(value) else None,
                        "published_text": str(value),
                        "source_row": source_row,
                        "source_column": column,
                    }
                )
    wb.close()
    igmh_profiles = []
    for metric in sorted({r["metric"] for r in igmh}):
        rows = [r for r in igmh if r["metric"] == metric]
        valid = [r for r in rows if r["value"] is not None]
        igmh_profiles.append(
            {
                "metric": metric,
                "reported_cells": len(rows),
                "numeric_cells": len(valid),
                "first_year": min(r["year"] for r in valid),
                "last_year": max(r["year"] for r in valid),
                "unavailable_years": [r["year"] for r in rows if r["value"] is None],
            }
        )
    igmh_by_year = defaultdict(dict)
    for r in igmh:
        igmh_by_year[r["year"]][r["metric"]] = r["value"]
    igmh_stay_checks = [
        {
            "year": year,
            "published_days": r["Average Duration of Stay (Days)"],
            "patient_days_per_reported_inpatient": r["In-patient Days"]
            / r["In-patients"],
            "definition_warning": "Diagnostic only: source definitions/cohorts may differ; never replace reported mean stay with this quotient.",
        }
        for year, r in sorted(igmh_by_year.items())
        if r.get("In-patients")
        and r.get("In-patient Days")
        and r.get("Average Duration of Stay (Days)") is not None
    ]
    with (OUT / "igmh-history-review.csv").open(
        "w", newline="", encoding="utf-8"
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=list(igmh[0]))
        writer.writeheader()
        writer.writerows(igmh)

    catalogue_directories = [
        "data/raw",
        "data/public/archive",
        "backend/data/cache/raw",
        "backend/data/cache/normalized",
        "backend/data/cache/derived",
        "backend/data/cache/graphs",
        "data/processed",
        "data/sample",
        "data/fixtures",
    ]
    directory_profiles = []
    for directory in catalogue_directories:
        files = [p for p in (ROOT / directory).rglob("*") if p.is_file()]
        directory_profiles.append(
            {
                "directory": directory,
                "files": len(files),
                "bytes": sum(p.stat().st_size for p in files),
                "extension_counts": dict(Counter(p.suffix or "none" for p in files)),
                "note": "File/cache counts are not independent dataset counts. Metadata, tiles, paginated source pages and derived graphs repeat lineage.",
            }
        )

    receipts_paths = [
        "backend/data/cache/normalized/islands.geojson",
        "backend/data/cache/normalized/population.json",
        "backend/data/cache/normalized/facilities.json",
        "backend/data/cache/normalized/transport.json",
        "backend/data/cache/normalized/provenance.json",
        "backend/data/cache/derived/accessibility-default.json",
        "backend/data/cache/derived/demand-anchors.json",
        "backend/data/cache/graphs/summary.json",
        "data/public/planning-baseline.json",
        "data/public/prediction-source.json",
        "data/public/context/agency-context.json",
        "data/public/context/followme-source.json",
        "data/public/policy-evidence/bundle.json",
        "data/public/policy-evidence/service-calendars.json",
        "data/public/operations/bundle.json",
        "data/public/delivery-evidence/references.json",
        "data/raw/census2022-population.xlsx",
        "data/raw/census2022-employment.xlsx",
        "data/raw/island-projections2022-2030.xlsx",
        "data/raw/population-definitions.xlsx",
        "data/raw/employment-definitions.xlsx",
        igmh_path,
        "data/public/archive/policy-evidence/school-statistics-2023.pdf",
        "data/public/archive/policy-evidence/constituency-population.xlsx",
        "data/public/archive/policy-evidence/constituency-employment.xlsx",
        "data/public/archive/policy-evidence/constituency-education.xlsx",
        "backend/data/cache/raw/mbs/P3.xlsx",
        "backend/data/cache/raw/mbs/P4.xlsx",
        "backend/data/cache/raw/mbs/P5.xlsx",
        "backend/data/cache/raw/mbs/P6.xlsx",
        "data/processed/demographics.json",
    ]

    source_records = []
    for path, rows in [
        ("data/public/planning-baseline.json", baseline["sources"]),
        ("data/public/policy-evidence/bundle.json", policy["sources"]),
        ("data/public/context/agency-context.json", context["sources"]),
        ("data/public/operations/bundle.json", operations["sources"]),
    ]:
        for source in rows:
            source_records.append(
                {
                    "artifact": path,
                    "id": source["id"],
                    "url": source.get("url"),
                    "hash": source.get("checksum", source.get("sha256")),
                    "reference_period": source.get(
                        "referencePeriod", source.get("referenceYear")
                    ),
                    "retrieved_at": source.get("retrievedAt", source.get("fetchedAt")),
                    "status": source.get("status"),
                }
            )

    inventory = {
        "generated_at": datetime.now(UTC).isoformat(),
        "scope": "All principal local source families: national normalized backend, Greater Male planner, policy/context/operations/public prediction archives; demo inputs excluded from analytical claims.",
        "method": "Offline read-only JSON and workbook profiling; exact administrative-atoll/name joins using current production reconciliation code; hashes freeze inspected originals. No online refresh, causal model or new prediction is claimed.",
        "source_receipts": [receipt(path) for path in receipts_paths],
        "directory_inventory": directory_profiles,
        "lineage": {
            "source_records_across_bundles": len(source_records),
            "distinct_id_checksum_pairs": len(
                {(r["id"], r["hash"]) for r in source_records}
            ),
            "records": source_records,
            "warning": "Different source IDs can still describe the same table; HDC page IDs are partitions of four layers. Facilities and destinations are the same45 records, not90 sites. Repeated source totals/sex/nationality/age rows are not additive.",
        },
        "national": {
            "islands": {
                **unique_check([x["properties"] for x in islands]),
                "category_counts": counts(
                    [x["properties"] for x in islands], "category"
                ),
            },
            "population": {
                "stored_rows": len(population),
                "source_counts": counts(population, "source"),
                "stored_match_status": counts(population_stored, "match_status"),
                "runtime_reconciled_match_status": counts(population, "match_status"),
                "runtime_key_check": unique_check(population),
                "matched_islands": len(island_census),
                "matched_residential_islands": sum(
                    island_properties[k].get("category") == "Residential Island"
                    for k in island_census
                ),
                "nonoverlapping_matched_residents_2022": sum(
                    x["total"] or 0 for x in island_census.values()
                ),
                "national_p4_republic_residents_2022": national_census_total,
                "national_p4_source_cell": "P4!C7",
                "national_residents_unallocated_to_current_island_summaries": national_census_total
                - sum(x["total"] or 0 for x in island_census.values()),
                "qualification": "190 islands are189 residential plus Hulhule; does not cover all national residents or residents on tourism/industrial islands. P6 atoll/city age counts remain unallocated. Raw 5,442 rows must never be summed. Stored metadata's1,049 attached records is superseded by current runtime reconciliation.",
                "greater_male": {
                    k: island_census[k]
                    for k in [
                        "island:1443",
                        "island:1442",
                        "island:1382",
                        "island:1483",
                    ]
                },
            },
            "facilities": {
                **unique_check(facilities),
                "categories": counts(facilities, "category"),
                "match_status": counts(facilities, "match_status"),
                "geocoded_records": sum(bool(x.get("geometry")) for x in facilities),
                "unlocated_records": sum(not x.get("geometry") for x in facilities),
                "island_context_records": sum(
                    bool(x.get("island_id")) for x in facilities
                ),
                "status_counts": dict(
                    Counter(
                        str(x.get("attributes", {}).get("status")) for x in facilities
                    )
                ),
                "by_category": [
                    {
                        "category": c,
                        "records": len(rows),
                        "geocoded": sum(bool(x.get("geometry")) for x in rows),
                        "with_island_id": sum(bool(x.get("island_id")) for x in rows),
                    }
                    for c, rows in [
                        (c, [r for r in facilities if r["category"] == c])
                        for c in sorted({x["category"] for x in facilities})
                    ]
                ],
                "duplicate_osm_claims": {
                    k: v
                    for k, v in (
                        (
                            k,
                            [
                                r["id"]
                                for r in facilities
                                if r.get("geometry") and r.get("osm_object_id") == k
                            ],
                        )
                        for k in {
                            r.get("osm_object_id")
                            for r in facilities
                            if r.get("geometry") and r.get("osm_object_id")
                        }
                    )
                    if len(v) > 1
                },
                "warning": "Registry rows do not establish current care availability, operating hours, seats or unique mapped campuses. Same OSM object claims can be distinct institutions at one site or unresolved duplicates; never equate site count with registry count.",
            },
            "transport": {
                **unique_check(transport["features"]),
                "categories": transport["summary"]["feature_categories"],
                "unassigned_to_island": sum(
                    not r.get("island_id") for r in transport["features"]
                ),
                "orphan_island_ids": sum(
                    bool(r.get("island_id")) and r["island_id"] not in island_ids
                    for r in transport["features"]
                ),
                "warning": "659 OSM features include hospitals/schools/government as well as transport and are not659 service stops.",
            },
        },
        "unused_island_indicators": {
            "population_workbook": {
                "source": receipt("data/raw/census2022-population.xlsx"),
                "administrative_island_rows": len(pop_indicator),
                "unique_atoll_locality_keys": len(
                    {
                        (
                            normalize_atoll_code(x["atoll_code"]),
                            normalize_name(x["locality"]),
                        )
                        for x in pop_indicator
                    }
                ),
                "safe_matched_islands": sum(
                    bool(x.get("island_id")) for x in pop_indicator
                ),
                "resident_2022_total": sum(x["values"][0] for x in pop_indicator),
                "maldivian_2022_total": sum(x["values"][3] for x in pop_indicator),
                "numeric_gaps": population_numeric_gaps,
                "reconciliation_failures": pop_reconciliation_failures,
                "cohort_definition": "Age/sex/dependency indicators apply to resident Maldivians, while total population and resident-density apply to all residents; excluded Greater Male and non-administrative islands cannot be imputed from this sheet.",
            },
            "employment_workbook": {
                "source": receipt("data/raw/census2022-employment.xlsx"),
                "administrative_island_rows": len(employment),
                "safe_matched_islands": sum(
                    bool(x.get("island_id")) for x in employment
                ),
                "unmatched": [
                    {
                        "source_row": r["source_row"],
                        "atoll": r["atoll_code"],
                        "locality": r["locality"],
                    }
                    for r in employment
                    if not r.get("island_id")
                ],
                "numeric_gaps": employment_numeric_gaps,
                "rates_outside_0_100": sum(
                    v < 0 or v > 100
                    for r in employment
                    for v in r["values"]
                    if finite_number(v)
                ),
                "cohort_definition": "Nine labour indicators by total/women/men, separately for resident total and resident Maldivians. Adult/labour-force/youth15–24/youth18–35 denominators differ. Published rates are percentages, not counts or transport trips. Weighted national rates require matching numerator/denominator counts.",
            },
        },
        "planning_baseline": {
            "reference_date": baseline["referenceDate"],
            "islands": len(baseline["islands"]),
            "origins": unique_check(baseline["origins"]),
            "origin_population_basis": "WorldPop2026alpha spatial weights normalized to official MBS2026/2030 island projections; modeled residence distribution, not actual passenger demand.",
            "population_by_island": [
                {
                    "id": x["id"],
                    "census2022": x.get("census2022Population"),
                    "projection2026": x.get("population2026"),
                    "projection2030": x.get("population2030"),
                    "origin_cells": sum(
                        r["islandId"] == x["id"] for r in baseline["origins"]
                    ),
                    "origin_population2026_sum": sum(
                        r["population2026"]
                        for r in baseline["origins"]
                        if r["islandId"] == x["id"]
                    ),
                }
                for x in baseline["islands"]
            ],
            "network_nodes": len(baseline["network"]["nodes"]),
            "network_edges": len(baseline["network"]["edges"]),
            "network_edges_missing_nodes": sum(
                r[k] not in baseline_node_ids
                for r in baseline["network"]["edges"]
                for k in ["from", "to"]
            ),
            "facilities": len(baseline["facilities"]),
            "destinations_duplicate_facilities": baseline["destinations"]
            == baseline["facilities"],
            "buildings": len(baseline["buildings"]["features"]),
            "stops": unique_check(baseline["transit"]["stops"]),
            "patterns": len(baseline["transit"]["patterns"]),
            "routes": len(baseline["transit"]["routes"]),
            "pattern_modes": counts(baseline["transit"]["patterns"], "mode"),
            "service_coverage": baseline["transit"]["serviceCoverage"],
            "cost_anchors": baseline["costAnchors"],
            "demographics": baseline["demographics"],
        },
        "policy": {
            "layers": layer_profiles,
            "aggregate_observations": len(policy["observations"]),
            "observations_by_source": counts(policy["observations"], "sourceId"),
            "constituencies": len(policy["constituencies"]),
            "constituencies_with_boundaries": sum(
                r["boundaryMatched"] for r in policy["constituencies"]
            ),
            "school_completeness": policy["schoolCompleteness"],
        },
        "igmh_history": {
            "source": receipt(igmh_path),
            "table_title": "IN-PATIENT HEALTH CARE IN IGMH2010–2025; older2005–2009 columns are populated for some metrics",
            "series": igmh_profiles,
            "currently_exposed_observations": sum(
                r["sourceId"] == "mbs-igmh-utilisation-2025"
                for r in policy["observations"]
            ),
            "stay_ratio_diagnostics": igmh_stay_checks,
            "limits": "Institutional activity rather than unique people or localized need. Year-end beds differ from average available beds. Reported utilization cannot predict current appointment availability or assign activity to a second Dharumavantha campus.",
        },
        "agency_context": {
            "vessels": {
                "port_call_rows": len(context["layers"]["vessels"]["portCalls"]),
                "quality_checks": context["layers"]["vessels"]["qualityChecks"],
                "live_positions_available": context["layers"]["vessels"][
                    "livePositionsAvailable"
                ],
            },
            "flights": {
                "observation_rows": len(context["layers"]["flights"]["observations"]),
                "quality_checks": context["layers"]["flights"]["qualityChecks"],
                "reference_period": context["layers"]["flights"]["referencePeriod"],
                "warning": "Monthly data span twelve months of2025 only;210mixed monthly/component/annual observations are not210months or independent samples.",
            },
            "followme": read("data/public/context/followme-source.json")["coverage"],
        },
        "operations": {
            "weather_alerts": len(operations["weatherAlerts"]),
            "alert_min_valid_from": min(
                r["validFrom"] for r in operations["weatherAlerts"]
            ),
            "alert_max_valid_until": max(
                r["validUntil"] for r in operations["weatherAlerts"]
            ),
            "road_events": len(operations["roadEvents"]),
            "notice_candidates": len(operations["noticeCandidates"]),
            "national_projection_readiness_rows": len(operations["nationalReadiness"]),
            "service_evidence": operations["serviceEvidence"],
            "warning": "Twenty latest RSS alerts are a truncated historical snapshot, not an exhaustive training archive. Road event absence does not prove uncongested roads.208readiness rows do not establish national routes/accessibility.",
        },
        "prediction": {
            "as_of": prediction["asOf"],
            "series": [
                {
                    "id": r["id"],
                    "points": len(r["points"]),
                    "last_actual_month": r["lastActualMonth"],
                    "unit": r["unit"],
                    "qualification": r["qualification"],
                }
                for r in prediction["series"]
            ],
            "audit": prediction["audit"],
        },
        "calendar_evidence": {
            "calendars": len(calendars["calendars"]),
            "routing_eligible": sum(
                r["routingEligible"] for r in calendars["calendars"]
            ),
            "with_crossing_minutes": sum(
                r["crossingMinutes"] is not None for r in calendars["calendars"]
            ),
            "qualification": "Four calendar documents have departures/day types but crossing durations or current validity are unresolved; cannot manufacture date-complete return journeys.",
        },
        "excluded_demo": {
            "processed_dataset": read("data/processed/demographics.json")["dataset"],
            "sample_and_fixtures": ["data/sample", "data/fixtures", "data/processed"],
            "warning": "Processed ingestion output comes from hand-made fixtures, not a second real demographic source.",
        },
    }
    (OUT / "inventory-review.json").write_text(
        json.dumps(inventory, indent=2, ensure_ascii=False, allow_nan=False),
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "output": str(OUT / "inventory-review.json"),
                "pop_indicator_rows": len(pop_indicator),
                "employment_matched": sum(bool(x.get("island_id")) for x in employment),
                "census_islands": len(island_census),
                "igmh_series": igmh_profiles,
            },
            ensure_ascii=True,
        )
    )


if __name__ == "__main__":
    main()
