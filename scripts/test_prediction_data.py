"""Network-free regression checks for publication-layout drift guards."""
import copy
import importlib.util
import unittest
from pathlib import Path

from openpyxl import Workbook

spec = importlib.util.spec_from_file_location("prediction_data", Path(__file__).with_name("prediction-data.py"))
source = importlib.util.module_from_spec(spec)
spec.loader.exec_module(source)


def worksheet():
    sheet = Workbook().active
    sheet.title = "June 2026"
    labels = ["Jan", "Feb", "Mar", "April", "May ", "June ", "July", "Aug", "Sep ", "Oct", "Nov", "Dec"]
    for offset in range(54):
        sheet.cell(4, 28 + offset, labels[offset % 12])
        if offset % 12 == 0:
            sheet.cell(3, 28 + offset, 2022 + offset // 12)
    sheet.cell(3, 30, 2022)  # A repeated source year is allowed only if correct.
    sheet.cell(161, 3, source.EXPECTED_LABELS["ferry_male_vm"])
    return sheet


def audited_fixture():
    return {
        "created_at_utc": "2026-10-11T00:00:00+00:00",
        "source_provenance": {key: {"url": url} for key, url in source.URLS.items()},
        "series": {
            key: {"source_row": row, "source_label": source.EXPECTED_LABELS[key],
                  "observations": [{"month": month, "value": 123} for month in source.expected_months(end - 27)],
                  "older_vintage_overlap": []}
            for key, _, _, _, _, row, end in source.CONFIG
        },
        "audit": {"june_2026_six_bus_series_repeat_may": [
            {"source_row": row, "label": label, "equal": True, "matches_prior_may_release": True}
            for row, label in source.EXPECTED_BUS_LABELS.items()
        ]},
    }


class PublicationGuards(unittest.TestCase):
    def test_audited_month_aliases_and_repeated_year(self):
        self.assertEqual(source.verify_calendar(worksheet(), 28, 54), source.expected_months(54))

    def test_shifted_month_header_stops_import(self):
        sheet = worksheet()
        sheet.cell(4, 29, "Jan")
        with self.assertRaisesRegex(ValueError, "Calendar header drift"):
            source.verify_calendar(sheet, 28, 54)

    def test_year_anchor_or_repeated_year_drift_stops_import(self):
        for column in (40, 30):
            sheet = worksheet()
            sheet.cell(3, column, 2021)
            with self.assertRaisesRegex(ValueError, "Calendar year drift"):
                source.verify_calendar(sheet, 28, 54)

    def test_changed_category_cannot_become_new_trusted_label(self):
        sheet = worksheet()
        sheet.cell(161, 3, "Male' ( HM Ferry )")
        with self.assertRaisesRegex(ValueError, "Category drift"):
            source.verify_category(sheet, 161, source.EXPECTED_LABELS["ferry_male_vm"])

    def test_existing_audit_path_checks_middle_dates_and_category(self):
        fixture = audited_fixture()
        self.assertEqual(len(source.prepare(fixture)["series"]), 3)
        changed = copy.deepcopy(fixture)
        changed["series"]["ferry_male_vm"]["observations"][10]["month"] = "2022-10"
        with self.assertRaisesRegex(ValueError, "Audit calendar range changed"):
            source.prepare(changed)
        changed = copy.deepcopy(fixture)
        changed["series"]["ferry_male_vm"]["source_label"] = "Wrong ferry"
        with self.assertRaisesRegex(ValueError, "Audit source category changed"):
            source.prepare(changed)

    def test_invalid_counts_and_changed_quarantine_labels_rejected(self):
        for value in (True, -1, 0.5, float("nan"), float("inf")):
            changed = audited_fixture()
            changed["series"]["bus_male_hulhumale"]["observations"][0]["value"] = value
            with self.assertRaisesRegex(ValueError, "Audit passenger counts invalid"):
                source.prepare(changed)
        changed = audited_fixture()
        changed["audit"]["june_2026_six_bus_series_repeat_may"][0]["label"] = "Different category"
        with self.assertRaisesRegex(ValueError, "Repeated-value audit categories changed"):
            source.prepare(changed)


if __name__ == "__main__":
    unittest.main()
