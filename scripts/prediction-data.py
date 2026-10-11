"""Prepare the public monthly prediction source from MBS publication worksheets.

Requirements: requests, openpyxl. Default downloads four official XLSX releases.
An existing independently audited JSON can be used with --audit PATH.
No internal Workings sheets or raw vehicle identifiers are published.
"""
import argparse
import hashlib
import io
import json
import math
from datetime import datetime, timezone, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
URLS = {
    "latest_june_2026": "https://statisticsmaldives.gov.mv/mbs/wp-content/uploads/2026/09/MIF_June-2026.xlsx",
    "may_2026": "https://statisticsmaldives.gov.mv/mbs/wp-content/uploads/2026/07/MIF-May-2026.xlsx",
    "dec_2023_issue": "https://statisticsmaldives.gov.mv/mbs/wp-content/uploads/2024/02/MIF-Dec-2023.xlsx",
    "dec_2024_issue": "https://statisticsmaldives.gov.mv/mbs/wp-content/uploads/2025/01/MIF_Dec-2024.xlsx",
}
CONFIG = [
    ("ferry_male_vm", "ferry-male-vm", "Malé · VM ferry endpoint", "ferry", "Malé endpoint of the published VM ferry category", 161, 81),
    ("ferry_villigili_vm", "ferry-villigili-vm", "Villigili · VM ferry endpoint", "ferry", "Villigili endpoint of the published VM ferry category", 162, 81),
    ("bus_male_hulhumale", "bus-male-hulhumale", "Malé / Hulhumalé bus movements", "bus", "Published Malé/Hulhumalé bus category; no route or stop allocation", 167, 80),
]
EXPECTED_LABELS = {
    "ferry_male_vm": "Male' ( VM Ferry )",
    "ferry_villigili_vm": "Villigili' ( VM Ferry )",
    "bus_male_hulhumale": "Male'/Hulhumale",
}
EXPECTED_BUS_LABELS = {
    166: "Airport", 167: "Male'/Hulhumale", 168: "Hulhumale Phase 1 & Phase2",
    169: "Male'", 170: "Villimale Internal", 171: "Free of Charge Passengers",
}
PUBLICATION_LAYOUTS = {
    "latest_june_2026": ("June 2026", 28, 54),
    "may_2026": ("May 2026", 28, 53),
    # This older release's first publication sheet has an anomalous title.
    "dec_2023_issue": ("Nov 2022", 28, 23),
    "dec_2024_issue": ("dec 2024", 62, 35),
}
MONTH_HEADERS = (
    {"jan", "january"}, {"feb", "february"}, {"mar", "march"},
    {"apr", "april"}, {"may"}, {"jun", "june"}, {"jul", "july"},
    {"aug", "august"}, {"sep", "sept", "september"}, {"oct", "october"},
    {"nov", "november"}, {"dec", "december"},
)


def expected_months(count):
    return [f"{2022 + i // 12}-{i % 12 + 1:02d}" for i in range(count)]


def verify_calendar(sheet, start, count):
    """Check publication dates before deriving any observation dates from columns."""
    months = expected_months(count)
    for offset, month in enumerate(months):
        column = start + offset
        label = sheet.cell(4, column).value
        if not isinstance(label, str) or label.strip().lower() not in MONTH_HEADERS[offset % 12]:
            raise ValueError(f"Calendar header drift: {sheet.title}, column {column}, expected {month}, got {label!r}")
        year = sheet.cell(3, column).value
        expected_year = 2022 + offset // 12
        # Check every January anchor and any explicitly repeated year cell. Blank
        # non-January cells are normal; never forward-fill an unchecked year.
        if offset % 12 == 0 or year is not None:
            if isinstance(year, bool) or not isinstance(year, (int, float)) or year != expected_year:
                raise ValueError(f"Calendar year drift: {sheet.title}, column {column}, expected {expected_year}, got {year!r}")
    return months


def verify_category(sheet, row, expected):
    actual = sheet.cell(row, 3).value
    if actual != expected:
        raise ValueError(f"Category drift: {sheet.title}, row {row}, expected {expected!r}, got {actual!r}")


def valid_count(value):
    return (isinstance(value, (int, float)) and not isinstance(value, bool)
            and math.isfinite(value) and value >= 0 and value == int(value))


def download_audit():
    import requests
    import openpyxl
    sheets, provenance = {}, {}
    for key, url in URLS.items():
        response = requests.get(url, timeout=90)
        response.raise_for_status()
        sheet = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True).worksheets[0]
        title, start, count = PUBLICATION_LAYOUTS[key]
        if sheet.title != title:
            raise ValueError(f"Publication worksheet changed: {key}, expected {title!r}, got {sheet.title!r}")
        verify_calendar(sheet, start, count)
        sheets[key] = sheet
        provenance[key] = {"url": url, "sha256": hashlib.sha256(response.content).hexdigest(), "publication_sheet": sheet.title}
    latest = sheets["latest_june_2026"]
    for row, heading in {158: "11.2 Male' area ferry passenger movement", 165: "11.3 Male' area Bus passenger movement"}.items():
        if latest.cell(row, 2).value != heading:
            raise ValueError(f"Source table heading drift: row {row}")
    for key, _, _, _, _, row, _ in CONFIG:
        verify_category(latest, row, EXPECTED_LABELS[key])
    repeated = []
    for row in range(166, 172):
        verify_category(latest, row, EXPECTED_BUS_LABELS[row])
        verify_category(sheets["may_2026"], row, EXPECTED_BUS_LABELS[row])
        may, june = latest.cell(row, 80).value, latest.cell(row, 81).value
        repeated.append({"source_row": row, "label": latest.cell(row, 3).value, "may": may, "june": june,
                         "equal": may == june, "matches_prior_may_release": may == sheets["may_2026"].cell(row, 80).value})
    if not all(item["equal"] and item["matches_prior_may_release"] for item in repeated):
        raise ValueError("June bus source changed: re-audit quarantine before publishing")
    result = {"created_at_utc": datetime.now(timezone.utc).isoformat(), "source_provenance": provenance,
              "audit": {"june_2026_six_bus_series_repeat_may": repeated}, "series": {}}
    for key, _, _, _, _, row, end in CONFIG:
        observations = []
        for column in range(28, end + 1):
            value = latest.cell(row, column).value
            if not valid_count(value):
                raise ValueError(f"Missing/invalid calendar observation: {key}, column {column}")
            observations.append({"month": f"{2022 + (column - 28) // 12}-{(column - 28) % 12 + 1:02d}", "value": value})
        label = EXPECTED_LABELS[key]
        checks = []
        for vintage, start, count in [("dec_2023_issue", 28, 23), ("dec_2024_issue", 62, 35), ("may_2026", 28, 53)]:
            sheet = sheets[vintage]
            old_rows = [r for r in range(1, sheet.max_row + 1) if sheet.cell(r, 3).value == label]
            if len(old_rows) != 1:
                raise ValueError(f"Category definition/row changed in {vintage}: {label}")
            differences = [{"month": observations[i]["month"], "older": sheet.cell(old_rows[0], start + i).value,
                            "latest": observations[i]["value"]} for i in range(count)
                           if sheet.cell(old_rows[0], start + i).value != observations[i]["value"]]
            checks.append({"vintage": vintage, "older_row": old_rows[0], "compared_months": count, "differences": differences})
        result["series"][key] = {"source_row": row, "source_label": label, "observations": observations,
                                  "older_vintage_overlap": checks, "last_trusted_month": observations[-1]["month"]}
    return result


def prepare(audit):
    series = []
    for key, public_id, name, mode, geography, row, end in CONFIG:
        record = audit["series"][key]
        points = record["observations"]
        expected_last = "2026-05" if mode == "bus" else "2026-06"
        if record["source_row"] != row or record["source_label"] != EXPECTED_LABELS[key]:
            raise ValueError(f"Audit source category changed: {key}")
        if [point["month"] for point in points] != expected_months(end - 27):
            raise ValueError(f"Audit calendar range changed: {key}")
        if not all(valid_count(point["value"]) for point in points):
            raise ValueError(f"Audit passenger counts invalid: {key}")
        series.append({"id": public_id, "name": name, "mode": mode, "geography": geography,
                       "lastActualMonth": expected_last, "unit": "passenger movements", "sourceUrl": URLS["latest_june_2026"],
                       "points": points, "qualification": f"Official source label: {record['source_label']}. Monthly passenger movements, not unique people. "
                       + ("June 2026 bus values quarantined as ambiguous carry-forward. No route/peak/stop demand allocation."
                          if mode == "bus" else "Endpoint direction/boarding versus alighting interpretation is undocumented; endpoints remain separate.")})
    repeat = audit["audit"]["june_2026_six_bus_series_repeat_may"]
    if len(repeat) != 6 or not all(x["equal"] and x["matches_prior_may_release"] for x in repeat):
        raise ValueError("Invalid repeated-value audit")
    if [(item["source_row"], item["label"]) for item in repeat] != list(EXPECTED_BUS_LABELS.items()):
        raise ValueError("Repeated-value audit categories changed")
    notes = [
        "Experimental aggregate outlook. Forecasts are not predicted riders, road congestion or capacity utilization.",
        "All six June 2026 bus categories repeat May exactly; quarantine June as ambiguous carry-forward, without claiming proven imputation.",
        "Ferry endpoint categories are kept separate. Boarding/alighting and directional definitions are unresolved.",
        "Free-of-charge passengers are not added to route categories because overlap is undocumented.",
        "Published history has been explored and uses revised historical values. Backtests are exploratory revised-vintage evaluation, not an untouched final or real-time publication-vintage test.",
        "Latest checked monthly issue June 2026 was published 7 September 2026. Last actual ferry month June; trusted bus month May. Intervening months are unobserved estimates.",
        "No explicit open redistribution licence has been established for the RTL coordinate capability. Vehicle capture stays private; only aggregate receipt is public.",
    ]
    return {"version": 1, "asOf": audit["created_at_utc"], "sourceIds": list(audit["source_provenance"]),
            "dataQualityNotes": notes, "series": series,
            "audit": {"sources": audit["source_provenance"], "publishedAt": "2026-09-07", "sourceTables": ["11.2", "11.3"],
                      "quarantinedBusMonth": "2026-06", "repeatedBusValues": repeat,
                      "crossVintage": {key: audit["series"][key]["older_vintage_overlap"] for key, *_ in CONFIG},
                      "revisionVintageBias": notes[4], "definitionStatus": "published passenger movement; finer interpretation unresolved",
                      "transformVersion": "mbs-monthly-source-v1"}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--audit", type=Path, help="Use an existing audited artifact instead of downloading")
    parser.add_argument("--output", type=Path, default=ROOT / "data/public/prediction-source.json")
    parser.add_argument("--capture-receipt", type=Path, help="Include only aggregate evidence from a bounded private RTL capture")
    args = parser.parse_args()
    audit = json.loads(args.audit.read_text(encoding="utf-8")) if args.audit else download_audit()
    result = prepare(audit)
    if args.capture_receipt:
        receipt = json.loads(args.capture_receipt.read_text(encoding="utf-8"))
        if receipt["sourceUrl"] != "https://bo.rtl.mv:4455/maldives/api/booking/v1/bus/livecoordinates" or receipt["sourceRequestRouteCode"] != "133" or receipt["measurementTimestampAvailable"] is not False:
            raise ValueError("Unexpected capture source, route or timestamp claim")
        quality = receipt["quality"]
        result["audit"]["monthlyAuditAsOf"] = result["asOf"]
        if datetime.fromisoformat(receipt["completedAt"].replace("Z", "+00:00")) > datetime.fromisoformat(result["asOf"].replace("Z", "+00:00")):
            result["asOf"] = receipt["completedAt"]
        days = {datetime.fromisoformat(x["receivedAt"].replace("Z", "+00:00")).astimezone(timezone(timedelta(hours=5))).date().isoformat() for x in quality}
        result["audit"]["busCapture"] = {
            "status": "completed", "sourceUrl": receipt["sourceUrl"], "routeId": "133", "startedAt": receipt["startedAt"],
            "completedAt": receipt["completedAt"], "snapshots": receipt["snapshots"],
            "vehicles": max((x["vehicleCount"] for x in quality), default=0), "observedDays": len(days),
            "changedPositionIntervals": sum(x["changedPositions"] > 0 for x in quality),
            "positionChanges": receipt["changedPositionComparisons"], "measurementTimestampAvailable": False,
            "firstReceivedAt": quality[0]["receivedAt"] if quality else None,
            "lastReceivedAt": quality[-1]["receivedAt"] if quality else None,
            "stopReason": receipt["stopReason"], "intervalSeconds": receipt["options"]["intervalSeconds"],
            "redistributionLicence": "Unverified; normalized positions retained in ignored private research archive",
            "qualification": "Bounded sample collected; collector stopped. No validated bus forecast available. " + receipt["qualification"],
        }
        result["dataQualityNotes"].append("A bounded route-133 capture verified changing public coordinates. Only aggregate capture evidence is published; no source measurement timestamp or evaluated passage-time forecast exists.")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), "series": [{"id": x["id"], "months": len(x["points"]), "lastActualMonth": x["lastActualMonth"]} for x in result["series"]]}))


if __name__ == "__main__":
    main()
