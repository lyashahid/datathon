"""Deterministic preparation of cached Greater Malé data; never invent inputs."""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, timezone, timedelta
import hashlib
import heapq
import json
import math
from pathlib import Path
import re
import unicodedata

import numpy as np
import openpyxl
import osmium
import pdfplumber
import rasterio
from bs4 import BeautifulSoup
from shapely.geometry import Point, LineString, Polygon, box, mapping, shape
from shapely.ops import polygonize, unary_union


def now():
    return datetime.now(timezone.utc).isoformat()


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":"), allow_nan=False), encoding="utf-8")


def norm(value):
    s = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]", "", s.lower())


def distance(a, b):
    lon1, lat1, lon2, lat2 = map(math.radians, [a[0], a[1], b[0], b[1]])
    x = math.sin((lat2-lat1)/2)**2 + math.cos(lat1)*math.cos(lat2)*math.sin((lon2-lon1)/2)**2
    return 6371000 * 2 * math.asin(min(1, math.sqrt(x)))


def coords_round(coords):
    return [[round(x, 7), round(y, 7)] for x, y in coords]


class OsmSnapshot(osmium.SimpleHandler):
    def __init__(self, bounds):
        super().__init__()
        self.bounds = bounds
        self.coast = []
        self.roads = []
        self.pois = []
        self.buildings = []
        self.bus_stops = []
        self.timestamp = None

    def contains(self, lon, lat):
        a, b, c, d = self.bounds
        return a <= lon <= c and b <= lat <= d

    def node(self, n):
        if not n.location.valid() or not self.contains(n.location.lon, n.location.lat):
            return
        tags = dict(n.tags)
        coord = (n.location.lon, n.location.lat)
        if tags.get("amenity") in {"school", "hospital", "clinic", "college", "university"} or tags.get("healthcare") in {"hospital", "clinic"}:
            self.pois.append(dict(osmType="node", osmId=n.id, tags=tags, coordinate=coord))
        if tags.get("highway") == "bus_stop" or tags.get("amenity") == "ferry_terminal":
            self.bus_stops.append(dict(osmType="node", osmId=n.id, tags=tags, coordinate=coord))

    def way(self, w):
        try:
            coordinates = [(n.lon, n.lat) for n in w.nodes]
        except (osmium.InvalidLocationError, RuntimeError):
            return
        if len(coordinates) < 2 or not any(self.contains(*c) for c in coordinates):
            return
        tags = dict(w.tags)
        if tags.get("natural") == "coastline":
            self.coast.append(LineString(coordinates))
        highway = tags.get("highway")
        if highway and highway not in {"construction", "proposed", "raceway", "bus_stop", "platform", "elevator"}:
            self.roads.append(dict(id=w.id, tags=tags, nodes=[n.ref for n in w.nodes], coordinates=coordinates))
        if tags.get("amenity") in {"school", "hospital", "clinic", "college", "university"} or tags.get("healthcare") in {"hospital", "clinic"}:
            geom = Polygon(coordinates) if coordinates[0] == coordinates[-1] and len(coordinates) >= 4 else LineString(coordinates)
            p = geom.representative_point()
            self.pois.append(dict(osmType="way", osmId=w.id, tags=tags, coordinate=(p.x, p.y)))
        if tags.get("building") and coordinates[0] == coordinates[-1] and len(coordinates) >= 4:
            geom = Polygon(coordinates)
            if geom.is_valid:
                self.buildings.append(dict(id=w.id, tags=tags, geometry=geom))
        if tags.get("amenity") == "ferry_terminal":
            p = LineString(coordinates).centroid
            self.bus_stops.append(dict(osmType="way", osmId=w.id, tags=tags, coordinate=(p.x, p.y)))


def build_islands(osm, raw):
    polygons = list(polygonize(unary_union(osm.coast)))
    def at(coord):
        hits = [p for p in polygons if p.covers(Point(coord))]
        if not hits:
            raise ValueError(f"OSM coastline polygon not found at {coord}; no synthetic boundary fallback")
        return min(hits, key=lambda p: p.area)
    male = at((73.51, 4.175))
    vili = at((73.484, 4.175))
    connected = at((73.543, 4.226))
    # Airport and Hulhumalé are connected by a causeway. This spatial partition
    # keeps official Hulhumalé totals off the runway without claiming a surveyed
    # administrative boundary. The separating line cuts the narrow causeway.
    hulhumale = connected.intersection(box(73, 4.203, 74, 5))
    hulhule = connected.intersection(box(73, 4, 74, 4.203))
    geometry = {"male": male, "vilimale": vili, "hulhumale": hulhumale, "hulhule": hulhule}
    census = openpyxl.load_workbook(raw / "yearbook2025-table3.3.xlsx", data_only=True).active
    census_rows = {norm(row[2]): row for row in census.values if len(row) > 3 and isinstance(row[2], str) and isinstance(row[3], (int, float))}
    projection = openpyxl.load_workbook(raw / "island-projections2022-2030.xlsx", data_only=True)["Resident Population"]
    projection_rows = {norm(row[1]): row for row in projection.values if len(row) > 26 and isinstance(row[1], str) and isinstance(row[14], (int, float))}
    wards = ["henveiru", "galolhu", "machchangolhi", "maafannu"]
    populations = {
        "male": (sum(census_rows[k][3] for k in wards), sum(projection_rows[k][14] for k in wards), sum(projection_rows[k][26] for k in wards)),
        "vilimale": (census_rows["vilimaale"][3], projection_rows["villigili"][14], projection_rows["villigili"][26]),
        "hulhumale": (census_rows["hulhumaalephase1"][3] + census_rows["hulhumaalephase2"][3], projection_rows["hulhumaale"][14], projection_rows["hulhumaale"][26]),
        "hulhule": (census_rows["hulhule"][3], None, None),
    }
    names = {"male": "Malé", "vilimale": "Vilimalé", "hulhumale": "Hulhumalé", "hulhule": "Hulhulé"}
    islands = []
    for island_id, geom in geometry.items():
        if geom.is_empty or not geom.is_valid:
            raise ValueError(f"Invalid coastline geometry: {island_id}")
        pop22, pop26, pop30 = populations[island_id]
        islands.append(dict(id=island_id, name=names[island_id], geometry=mapping(geom.simplify(0.000008, preserve_topology=True)), bounds=list(geom.bounds), census2022Population=pop22, population2026=pop26, population2030=pop30, resident2026=pop26, resident2030=pop30, projection2026Population=pop26, projection2030Population=pop30, sourceIds=["osm-geofabrik", "mbs-census2022-male-table", "mbs-island-projections"], qualification="2026/2030 are official MBS projections, not observed residents. Harbours are excluded. Hulhulé 2026/2030 is unavailable because the official projection combines it with Thilafushi and Gulhifalhu." if island_id == "hulhule" else "2026/2030 are official MBS projections, not observed residents; calibrated projection baseline differs from the published 2022 census. Harbour population excluded. Hulhumalé/airport coastline partition is a derived causeway split, not an administrative boundary."))
    return islands, geometry


def locate_island(coord, geoms):
    p = Point(coord)
    for island_id, geom in geoms.items():
        if geom.covers(p):
            return island_id
    # Facility entrances and official stops can lie just outside coastline at a
    # quay. Assignment is qualified, and does not move the actual source point.
    nearest = min(geoms, key=lambda k: geoms[k].distance(p))
    return nearest if geoms[nearest].distance(p) < 0.0004 else None


def make_demographics(islands, raw):
    """Official island aggregates only; never assign a cohort to a grid cell."""
    census = openpyxl.load_workbook(raw/"yearbook2025-table3.3.xlsx",data_only=True).active
    observed = {norm(r[2]):r for r in census.values if len(r)>11 and isinstance(r[2],str) and isinstance(r[3],(int,float))}
    workbook = openpyxl.load_workbook(raw/"island-projections2022-2030.xlsx",data_only=True)
    projected = {sheet:{norm(r[1]):r for r in workbook[sheet].values if len(r)>28 and isinstance(r[1],str) and isinstance(r[14],(int,float))} for sheet in ["Resident Population","Maldivian Population","Resident Foreign population"]}
    observed_keys = dict(male=["henveiru","galolhu","machchangolhi","maafannu"],vilimale=["vilimaale"],hulhumale=["hulhumaalephase1","hulhumaalephase2"],hulhule=["hulhule"])
    projected_keys = dict(male=["henveiru","galolhu","machchangolhi","maafannu"],vilimale=["villigili"],hulhumale=["hulhumaale"])
    compiled = []
    for island in islands:
        iid = island["id"]
        rows = [observed[key] for key in observed_keys[iid]]
        denominator = sum(r[3] for r in rows)
        def group(dimension,year,status,denom,values,source):
            return dict(id=f"{iid}:{dimension}:{year}",dimension=dimension,year=year,referencePeriod=year,status=status,unit="people",denominator=denom,spatialGrain="physical island aggregate",groups=[dict(id=key,label=label,value=value) for key,label,value in values],sourceIds=[source],canAllocateToOrigins=False,qualification="Official aggregate for the physical island (Malé wards / Hulhumalé phases combined; harbour residents excluded). Projected values are not observed residents. No local cohort distribution or cohort-specific travel demand is supplied.")
        available = [group("sex",2022,"observed",denominator,[("female","Female",sum(r[4] for r in rows)),("male","Male",sum(r[5] for r in rows))],"mbs-census2022-male-table"),group("nationality",2022,"observed",denominator,[("maldivian","Maldivian residents",sum(r[6] for r in rows)),("foreign","Foreign residents",sum(r[9] for r in rows))],"mbs-census2022-male-table")]
        unavailable = [dict(dimension="age",status="unavailable",sourceIds=["mbs-census2022-population","mbs-population-definitions"],reason="Acquired age indicators cover administrative islands in atolls, excluding Malé localities. No matched age group at Greater Malé island or grid grain is supplied; national/atoll values are not allocated locally."),dict(dimension="disability",status="unavailable",sourceIds=["mbs-census2022-population","mbs-population-definitions"],reason="No disability cohort at Greater Malé island or grid grain is present in the acquired population sources. Accessibility tags on institutions do not replace a resident disability count."),dict(dimension="employment",status="unavailable",sourceIds=["mbs-census2022-employment","mbs-employment-definitions"],reason="Acquired labour-force island indicators cover administrative islands in atolls, excluding Malé localities. No matched Greater Malé resident employed/job-seeker distribution or workplace inventory is supplied.")]
        if iid in projected_keys:
            keys = projected_keys[iid]
            resident_rows = [projected["Resident Population"][key] for key in keys]
            for year,column in [(2026,14),(2030,26)]:
                denom = sum(r[column] for r in resident_rows)
                available.append(group("sex",year,"projected",denom,[("female","Female",sum(r[column+2] for r in resident_rows)),("male","Male",sum(r[column+1] for r in resident_rows))],"mbs-island-projections"))
                nationality = [("maldivian","Maldivian residents",sum(projected["Maldivian Population"][key][column] for key in keys)),("foreign","Foreign residents",sum(projected["Resident Foreign population"][key][column] for key in keys))]
                available.append(group("nationality",year,"projected",denom,nationality,"mbs-island-projections"))
        else:
            unavailable += [dict(dimension=dimension,years=[2026,2030],status="unavailable",sourceIds=["mbs-island-projections"],reason="Official projection combines Hulhulé with Thilafushi/Gulhifalhu; airport-island total and sex/nationality groups cannot be independently disaggregated.") for dimension in ["sex","nationality"]]
        record = dict(islandId=iid,spatialGrain="physical island aggregate",originCohortsAvailable=False,availableGroups=available,unavailableGroups=unavailable)
        compiled.append(record)
        island["demographics"] = record
    return dict(spatialGrain="physical island aggregate",originCohortsAvailable=False,islands=compiled,gaps=["Sex and nationality totals support an evidence drawer, not demographic heatmaps or demographic-specific access estimates. No neighborhood distribution is silently derived by applying uniform island ratios.","Age, disability and employment at valid Greater Malé local grain require further sources. Acquired administrative-atoll island workbooks exclude the study area."])


def make_origins(islands, geoms, raw):
    origins = []
    report = []
    with rasterio.open(raw / "worldpop2026.tif") as raster:
        a = raster.read(1)
        for island in islands:
            if island["population2026"] is None:
                continue
            geom = geoms[island["id"]]
            cells = []
            west, south, east, north = geom.bounds
            row0, col0 = raster.index(west, north)
            row1, col1 = raster.index(east, south)
            for r in range(max(0, row0-1), min(a.shape[0], row1+2)):
                for c in range(max(0, col0-1), min(a.shape[1], col1+2)):
                    value = float(a[r, c])
                    if not math.isfinite(value) or value <= 0 or value == raster.nodata:
                        continue
                    x1, y1 = raster.xy(r, c, offset="ul")
                    x2, y2 = raster.xy(r, c, offset="lr")
                    footprint = box(x1, y2, x2, y1).intersection(geom)
                    if footprint.is_empty or footprint.area <= 1e-12:
                        continue
                    # Boundary pixels are weighted by their land fraction. A
                    # representative point is guaranteed on the clipped land.
                    weight = value * footprint.area / abs((x2-x1)*(y2-y1))
                    point = footprint.representative_point()
                    cells.append((r, c, weight, footprint, (point.x, point.y)))
            denominator = sum(cell[2] for cell in cells)
            if denominator <= 0:
                raise ValueError(f"No positive WorldPop land pixels for {island['id']}; do not substitute uniform/synthetic origins")
            for r, c, weight, footprint, coordinate in cells:
                lon, lat = round(coordinate[0], 7), round(coordinate[1], 7)
                origins.append(dict(id=f"wp:{island['id']}:{r}:{c}", islandId=island["id"], lon=lon, lat=lat, coordinate=[lon, lat], bounds=list(footprint.bounds), geometry=mapping(footprint.simplify(0.000002, preserve_topology=True)), population2026=island["population2026"]*weight/denominator, population2030=island["population2030"]*weight/denominator, rasterWeight=weight, sourceIds=["worldpop2026", "mbs-island-projections", "osm-geofabrik"], provenance="simulated", qualification="WorldPop alpha spatial estimate normalized to official island projection; 2030 keeps 2026 spatial weights. Not block-level census or measured travel demand."))
            report.append(dict(islandId=island["id"], cells=len(cells), rawLandWeightedRasterPopulation=denominator, officialPopulation2026=island["population2026"], scaleFactor2026=island["population2026"]/denominator))
    return origins, report


def road_permissions(tags):
    h = tags.get("highway", "")
    forbidden = tags.get("access") in {"private", "no"}
    walk = h not in {"motorway", "motorway_link", "trunk", "trunk_link", "cycleway"} and not forbidden
    if tags.get("foot") in {"yes", "designated", "permissive"}:
        walk = True
    if tags.get("foot") in {"no", "private"}:
        walk = False
    drive = h not in {"footway", "pedestrian", "path", "steps", "cycleway", "corridor", "bridleway"} and not forbidden
    if tags.get("bus") in {"yes", "designated", "permissive"}:
        drive = True
    if tags.get("motor_vehicle") in {"no", "private"} and tags.get("bus") not in {"yes", "designated"}:
        drive = False
    if tags.get("motorcar") == "no" and tags.get("bus") not in {"yes", "designated"}:
        drive = False
    return walk, drive


def make_network(osm, geoms):
    selected = []
    for road in osm.roads:
        walk, drive = road_permissions(road["tags"])
        if not (walk or drive):
            continue
        # Keep actual connecting bridges as tagged OSM ways, including offshore
        # nodes. All other disconnected resort roads are outside this study.
        if not any(locate_island(c, geoms) for c in road["coordinates"]) and road["tags"].get("bridge") != "yes":
            continue
        # Split long source segments only along their existing geometry. These
        # extra points improve attachment precision without connecting roads
        # that do not share an OSM junction or changing any access direction.
        refs, cc = [road["nodes"][0]], [road["coordinates"][0]]
        for j, (a, b) in enumerate(zip(road["coordinates"], road["coordinates"][1:])):
            count = max(1, math.ceil(distance(a, b)/70))
            for k in range(1, count):
                refs.append(f"way{road['id']}:segment{j}:fraction{k}of{count}")
                cc.append((a[0]+(b[0]-a[0])*k/count, a[1]+(b[1]-a[1])*k/count))
            refs.append(road["nodes"][j+1])
            cc.append(b)
        selected.append(dict(road, nodes=refs, coordinates=cc))
    uses = Counter(n for road in selected for n in set(road["nodes"]))
    nodes = {}
    edges = []
    for road in selected:
        refs, cc, tags = road["nodes"], road["coordinates"], road["tags"]
        walk, drive = road_permissions(tags)
        retain = {0, len(refs)-1} | {i for i, ref in enumerate(refs) if uses[ref] > 1}
        cumulative = 0
        for i in range(1, len(refs)):
            segment = distance(cc[i-1], cc[i])
            if cumulative+segment > 70 and i-1 not in retain:
                retain.add(i-1)
                cumulative = 0
            cumulative += segment
            if i in retain or cumulative >= 70:
                retain.add(i)
                cumulative = 0
        # Closed loops without intersections need one interior retained point.
        if refs[0] == refs[-1] and len(retain) < 3:
            retain.add(len(refs)//2)
        indices = sorted(retain)
        reverse = tags.get("oneway") == "-1"
        one_way = tags.get("oneway") in {"yes", "1", "true", "-1"} or (tags.get("junction") == "roundabout" and tags.get("oneway") != "no")
        for start, end in zip(indices, indices[1:]):
            path = cc[start:end+1]
            if len(path) < 2:
                continue
            from_id, to_id = "osm:"+str(refs[start]), "osm:"+str(refs[end])
            if from_id == to_id:
                continue
            for node_id, coord in [(from_id, path[0]), (to_id, path[-1])]:
                nodes[node_id] = dict(id=node_id, lon=round(coord[0], 7), lat=round(coord[1], 7), coordinate=[round(coord[0], 7), round(coord[1], 7)], islandId=locate_island(coord, geoms), provenance="interpolated along OSM source segment" if ":segment" in node_id else "OSM node")
            if reverse:
                from_id, to_id, path = to_id, from_id, list(reversed(path))
            edges.append(dict(id=f"osmway:{road['id']}:{start}:{end}", from_=from_id, to=to_id, distanceM=round(sum(distance(a,b) for a,b in zip(path,path[1:])), 3), walkAllowed=walk, driveAllowed=drive, oneway=one_way, sourceOsmWayId=str(road["id"]), highway=tags.get("highway"), name=tags.get("name:en") or tags.get("name"), geometry=coords_round(path), bridge=tags.get("bridge") == "yes"))
            edges[-1]["from"] = edges[-1].pop("from_")
    return dict(nodes=list(nodes.values()), edges=edges, qualification="OSM node identity preserves junction topology. Degree-two geometry nodes are contracted into polylines; segments are split on their original geometry at a maximum 70 m spacing for attachment precision. Interpolated nodes introduce no new cross-road connections. Walking ignores motor oneway; driving follows it. No ferry crossings or proximity-based road connectors invented. Turn restrictions, pavement condition, signal delay and sidewalk quality are not modeled.")


def make_facilities(osm, geoms, raw, available):
    school_names = set()
    if "education-schools" in available:
        soup = BeautifulSoup((raw/"education-schools.html").read_text(encoding="utf-8"), "html.parser")
        school_names = {norm(x.get_text(" ", strip=True)) for x in soup.select("span.font-bold") if len(norm(x.get_text())) > 5}
    health_names = defaultdict(list)
    if "health-government2024" in available:
        with pdfplumber.open(raw/"government-health-facilities2024.pdf") as doc:
            for page in doc.pages:
                for line in (page.extract_text() or "").splitlines():
                    match = re.match(r"^\d+\s+(.*?)\s+GOV\s+(\S+)\s+HF-", line)
                    if match:
                        health_names[norm(match[1])].append(dict(name=match[1], institutionClass=match[2], sourceId="health-government2024"))
    for file in ["health-registry.html", "health-registry-hospitals.html", "health-registry-hospitals2.html", "health-registry-hospitals3.html"]:
        if Path(file).stem not in available:
            continue
        soup = BeautifulSoup((raw/file).read_text(encoding="utf-8"), "html.parser")
        for row in soup.select("tbody tr"):
            cols = [td.get_text(" ", strip=True) for td in row.select("td")]
            if len(cols) >= 3:
                health_names[norm(cols[0])].append(dict(name=cols[0], institutionClass=cols[2], sourceId=Path(file).stem))
    # These two documented spelling variants refer to the same named hospitals;
    # all other joins require an exact normalized institution name, not substring.
    aliases = {"indiragandhimemorialhospital": "indhiragandhimemorialhospital", "villimalehospital": "vilimalehospital"}
    facilities = []
    exclusions = []
    seen = set()
    for poi in osm.pois:
        tags, coordinate = poi["tags"], poi["coordinate"]
        island_id = locate_island(coordinate, geoms)
        name = tags.get("name:en") or tags.get("name")
        if not name or not island_id:
            continue
        name_normal = norm(name)
        if name_normal in seen:
            continue
        if name_normal in {"lobbya", "lobbyb"}:
            exclusions.append(dict(osmId=poi["osmId"], reason="Hospital lobby point is not a standalone destination"))
            continue
        seen.add(name_normal)
        kind = tags.get("amenity") or tags.get("healthcare")
        if kind == "university":
            kind = "college"
        if kind not in {"hospital", "clinic", "school", "college"}:
            continue
        sources = ["osm-geofabrik"]
        matched_health = health_names.get(aliases.get(name_normal, name_normal), []) if kind in {"hospital", "clinic"} else []
        official = (kind == "school" and name_normal in school_names) or bool(matched_health)
        if kind == "school" and official:
            sources.append("education-schools")
        if matched_health:
            sources += sorted({row["sourceId"] for row in matched_health})
            # Official categories such as General Clinic and Area Specific
            # Hospital must not be inferred from unrelated page text.
            category = matched_health[-1]["institutionClass"].lower()
            if "clinic" in category:
                kind = "clinic"
            elif "hospital" in category or category in {"t", "hg1", "hg2", "hg3"}:
                kind = "hospital"
        restricted = "military" in name_normal or tags.get("healthcare:speciality") == "addiction" or tags.get("access") in {"private", "no"}
        facilities.append(dict(id=f"osm:{poi['osmType']}:{poi['osmId']}", name=name, kind=kind, category=kind, islandId=island_id, lon=round(coordinate[0], 7), lat=round(coordinate[1], 7), coordinate=[round(coordinate[0], 7), round(coordinate[1], 7)], sourceIds=sources, officialNameMatches=matched_health, osmId=str(poi["osmId"]), osmType=poi["osmType"], verification=False if restricted else "official" if official else "osm-mapped", checkDate=tags.get("check_date"), wheelchair=tags.get("wheelchair"), emergency=tags.get("emergency"), healthcareSpeciality=tags.get("healthcare:speciality"), qualification="Restricted/specialist institution; excluded from general public destination analysis pending suitability confirmation." if restricted else "Exact normalized official directory name matched (IGMH has a documented spelling alias); coordinates are OSM-derived and services/capacity/public suitability are not audited." if official else "Mapped in OpenStreetMap; official list match not confirmed. Location/category completeness varies, capacity and general-public suitability are not audited."))
    exclusions.append(dict(officialName="Vilimale' Hospital", sourceIds=["health-registry-hospitals", "health-government2024"], reason="Official institution has no matching geocoded OSM object in snapshot. Not assigned an invented island-centroid coordinate; Vilimalé hospital accessibility is incomplete."))
    return facilities, exclusions


def minutes(value):
    h, m, *sec = value.split(":")
    return int(h)*60+int(m)+(int(sec[0])/60 if sec else 0)


def road_router(network):
    """Shortest directed road paths; geometry is retained from source ways."""
    nodes = {n["id"]: n for n in network["nodes"]}
    adjacency = defaultdict(list)
    for e in network["edges"]:
        if not e["driveAllowed"]:
            continue
        adjacency[e["from"]].append((e["to"], e["distanceM"], e["geometry"]))
        if not e["oneway"]:
            adjacency[e["to"]].append((e["from"], e["distanceM"], list(reversed(e["geometry"]))))
    drive_ids = set(adjacency) | {e[0] for edges in adjacency.values() for e in edges}
    snapped = {}
    solved = {}
    def attach(stop):
        if stop["id"] not in snapped:
            best = min(drive_ids, key=lambda node: distance(nodes[node]["coordinate"], stop["coordinate"]))
            offset = distance(nodes[best]["coordinate"], stop["coordinate"])
            snapped[stop["id"]] = (best, offset) if offset <= 100 else (None, offset)
        return snapped[stop["id"]]
    def path(a, b):
        start, a_offset = attach(a)
        end, b_offset = attach(b)
        if start is None or end is None:
            return None
        key = (start, end)
        if key not in solved:
            queue = [(0, start)]
            costs, previous = {start: 0}, {}
            while queue:
                cost, node = heapq.heappop(queue)
                if cost != costs.get(node):
                    continue
                if node == end:
                    break
                for nxt, length, geometry in adjacency[node]:
                    total = cost+length
                    if total < costs.get(nxt, math.inf):
                        costs[nxt] = total
                        previous[nxt] = (node, geometry)
                        heapq.heappush(queue, (total, nxt))
            if end not in costs:
                solved[key] = None
            else:
                pieces, cursor = [], end
                while cursor != start:
                    predecessor, geometry = previous[cursor]
                    pieces.append(geometry)
                    cursor = predecessor
                geometry = []
                for piece in reversed(pieces):
                    geometry += piece if not geometry else piece[1:]
                if not geometry:
                    geometry = [nodes[start]["coordinate"]]
                solved[key] = dict(distanceM=costs[end], geometry=geometry)
        return dict(solved[key], attachmentOffsetsM=[round(a_offset, 2), round(b_offset, 2)]) if solved[key] else None
    return path


def make_transit(raw, geoms, sources, network):
    gaps = ["Public API provides remaining bus departures at fetch time, not a complete recurring timetable. Calendar is pinned to snapshot date; morning and other-day results are incomplete.", "Return directions are only those explicitly supplied by the operator. No reverse services invented.", "Ferry frequency PDFs lack verified sailing runtime/precise terminals; retained as source evidence, excluded from journey calculations."]
    route_record = next((s for s in sources if s["id"] == "rtl-bus-routes"), None)
    if not route_record or route_record["status"] != "downloaded":
        return dict(stops=[], patterns=[], routes=[], gaps=gaps+["RTL bus route source unavailable"], accessStatus="unavailable"), None
    stamp = datetime.fromisoformat(route_record["fetchedAt"])
    reference_date = (stamp + timedelta(hours=5)).date().isoformat()
    data = json.loads((raw/"rtl-bus-routes.json").read_text(encoding="utf-8"))
    stops = {}
    patterns, routes = [], []
    discarded_trips = 0
    route_path = road_router(network)
    for r in data.get("routeResponse", []):
        raw_stops = sorted(r["busRouteStopList"], key=lambda s: s["order"])
        if len(raw_stops) < 2:
            continue
        stop_ids = []
        for stop in raw_stops:
            coord = (float(stop["longitude"]), float(stop["latitude"]))
            sid = "rtl:stop:"+str(stop["code"])
            island_id = locate_island(coord, geoms)
            if not island_id:
                gaps.append(f"Operator stop {stop['name']} is outside or offshore study coastline; excluded route {r['routeNumber']} from calculation.")
                stop_ids = []
                break
            if sid in stops and distance(stops[sid]["coordinate"], coord) > 15:
                sid += ":"+str(stop["id"])
            stops[sid] = dict(id=sid, name=stop["name"].strip(), islandId=island_id, lon=coord[0], lat=coord[1], coordinate=list(coord), sourceIds=["rtl-bus-routes"], operatorCode=str(stop["code"]), verification="operator-published")
            if "dropoffonly" in norm(stop["name"]):
                stops[sid]["pickupAllowed"] = False
            stop_ids.append(sid)
        if not stop_ids:
            continue
        rid = "rtl:route:"+str(r["code"])
        lookup = [{t["order"]: minutes(t["timing"]) for t in s["timings"]} for s in raw_stops]
        common = set.intersection(*(set(x) for x in lookup)) if lookup else set()
        trips = []
        for order in sorted(common):
            times = [lookup[i][order] for i in range(len(lookup))]
            # Midnight wrapping is explicit in times; add a day only where an
            # actual decreasing midnight sequence is plausible.
            for i in range(1,len(times)):
                if times[i] < times[i-1] and times[i-1] > 1380 and times[i] < 120:
                    times[i] += 1440
            if any(b<a or b-a>120 for a,b in zip(times,times[1:])):
                discarded_trips += 1
                continue
            trips.append(dict(id=f"{rid}:trip:{order}", stopTimes=[dict(arrivalMinutes=t, departureMinutes=t) for t in times]))
        source_ids = ["rtl-bus-routes", "osm-geofabrik"]
        roadshape_file = raw/("rtl-roadshape-"+str(r["code"])+".json")
        roadshape = None
        if any(s["id"] == "rtl-roadshape-"+str(r["code"]) and s["status"] == "downloaded" for s in sources):
            payload = json.loads(roadshape_file.read_text(encoding="utf-8"))
            operator_shape = payload.get("roadShape")
            if operator_shape and operator_shape.get("type") == "FeatureCollection":
                roadshape = dict(type="FeatureCollection",features=[])
                for feature in operator_shape.get("features", []):
                    geometry = shape(feature["geometry"]).simplify(0.000012,preserve_topology=True)
                    roadshape["features"].append(dict(type="Feature",properties=dict(sourceId="rtl-roadshape-"+str(r["code"]),qualification="Operator published road shape; direction/loop semantics not verified. Not used as a direction's full vehicle-kilometres."),geometry=mapping(geometry)))
                source_ids.append("rtl-roadshape-"+str(r["code"]))
        pieces = [route_path(stops[a],stops[b]) for a,b in zip(stop_ids,stop_ids[1:])]
        missing = [i for i,piece in enumerate(pieces) if piece is None]
        route_geometry = []
        if not missing:
            for piece in pieces:
                route_geometry += piece["geometry"] if not route_geometry else piece["geometry"][1:]
        if missing:
            gaps.append(f"{r['routeNumber']}: {len(missing)} consecutive stop pairs cannot be routed on the directed OSM graph within 100 m attachment; direction road distance unavailable. Operator road-shape overlay remains source evidence.")
        route = dict(id=rid, name=r["name"], routeNumber=r["routeNumber"], stopIds=stop_ids, mode="bus", color=r.get("color"), fareMvr=r.get("fare"), sourceIds=source_ids, timetableStatus="remaining-snapshot" if trips else "unavailable", geometry=route_geometry, operatorRoadShape=roadshape, distanceM=round(sum(piece["distanceM"] for piece in pieces),3) if not missing else None, segmentDistanceM=[piece["distanceM"] if piece else None for piece in pieces], segmentGeometry=[piece["geometry"] if piece else [] for piece in pieces], geometryStatus="incomplete" if missing else "modeled road path", geometryQualification="Shortest directed OSM road paths through operator ordered stops, snapped within 100 m. Modeled geometry, not a measured bus trajectory; turn restrictions and exact driver routing unavailable. Published operator shape retained separately with unverified direction semantics.")
        routes.append(route)
        if trips:
            patterns.append(dict(id=rid+":published-direction", routeId=rid, name=r["routeNumber"]+" · "+r["name"], mode="bus", stopIds=stop_ids, calendar=dict(weekdays=[], includedDates=[reference_date]), timing=dict(kind="scheduled", trips=trips), provenance="observed", sourceIds=source_ids, geometry=route_geometry, distanceM=route["distanceM"], segmentDistanceM=route["segmentDistanceM"], segmentGeometry=route["segmentGeometry"], color=r.get("color"), fareMvr=r.get("fare"), qualification="Operator public route API remaining-departure snapshot. Service date pinned to fetch date in MVT; direction is explicitly supplied. Zero dwell is an arrival/departure equivalence in the published timing, not a measured dwell. Direction road geometry/distance is an OSM shortest-path estimate, not operator vehicle-km."))
        else:
            gaps.append(f"{r['routeNumber']} {r['name']}: published stops exist but no complete timings returned; route visible, journey service unavailable.")
    if discarded_trips:
        gaps.append(f"{discarded_trips} non-monotonic or implausible timing-order joins rejected.")
    window = [min(t["stopTimes"][0]["departureMinutes"] for p in patterns for t in p["timing"]["trips"]), max(t["stopTimes"][-1]["arrivalMinutes"] for p in patterns for t in p["timing"]["trips"])] if patterns else None
    coverage = dict(snapshotDate=reference_date, remainingDeparturesOnly=True,startMinutes=window[0] if window else None,endMinutes=window[1] if window else None,fullDay=False,recurringCalendar=False,qualification="At 08:00 this snapshot cannot establish a complete bus comparison. Loaded later departures are available on snapshot date only.")
    route_record["serviceCoverage"] = coverage
    route_record["referencePeriod"] = reference_date
    return dict(stops=list(stops.values()), patterns=patterns, routes=routes, gaps=gaps, accessStatus="partial operator snapshot", snapshotDate=reference_date, fetchedAt=route_record["fetchedAt"], recordedDepartureWindowMinutes=window,serviceCoverage=coverage), reference_date


def make_cost_anchors(raw, available):
    anchors = dict(dieselRetailMvrPerLitre=None,petrolRetailMvrPerLitre=None,electricityCommercialBands=None,gridKgCo2ePerKwh=None,qualification="No sourced marginal bus/stop capital cost, full per-km operating cost, or certified current Malé charging-grid factor. User-input illustrations must be clearly labeled.")
    if "sto-fuel2026" in available:
        text = BeautifulSoup((raw/"sto-fuel2026.html").read_text(encoding="utf-8"),"html.parser").get_text(" ",strip=True)
        for fuel in ["Diesel","Petrol"]:
            match = re.search(fuel+r" price.*?New price:\s*([\d.]+)", text)
            if match:
                anchors[fuel.lower()+"RetailMvrPerLitre"] = dict(value=float(match[1]),unit="MVR/litre",effectiveDate="2026-03-05",sourceIds=["sto-fuel2026"],qualification="Published STO retail revision. Operator contracted/wholesale price unverified; this is not total operating cost.")
    if "stelco-tariff" in available:
        soup = BeautifulSoup((raw/"stelco-tariff.html").read_text(encoding="utf-8"),"html.parser")
        table = next((t for t in soup.select("table") if t.get_text(" ",strip=True).startswith("# Band (30 Days) Commercial")),None)
        bands = []
        if table:
            for row in table.select("tr"):
                cells = [c.get_text(" ",strip=True) for c in row.select("td")]
                if len(cells) != 3:
                    continue
                limits = [int(v.replace(",","")) for v in re.findall(r"\d[\d,]*",cells[1])]
                if len(limits) == 2:
                    lower, upper = limits
                elif len(limits) == 1 and ">" in cells[1]:
                    lower, upper = limits[0]+1,None
                else:
                    continue
                bands.append(dict(minMonthlyKwh=lower,maxMonthlyKwh=upper,rateMvrPerKwh=float(cells[2])))
        if len(bands) == 7:
            anchors["electricityCommercialBands"] = dict(unit="MVR/kWh",sourceIds=["stelco-tariff"],bands=bands,qualification="Published Malé commercial tariff. Page has older surcharge notes; current depot category, surcharge and bill not verified.")
    return anchors


def audit_transit_stops(transit, network, geoms, raw, sources):
    """Audit catalogue identity, loaded service and location as separate grains.

    An operator-listed stop is not proof that it operates on every day. A missing
    remaining-departure timing is not proof that a physical stop has closed.
    """
    audit_sources = []
    current_data = None
    definitions = [
        ("rtl-routes-current", "rtl-bus-routes-audit", "RTL stop catalogue fresh audit snapshot", "https://bo.rtl.mv:4455/maldives/api/booking/v2/bus/routedetails", "Unauthenticated public operator endpoint; feed reuse license not stated"),
        ("mtcc-westpark-notice-report", "mtcc-westpark-notice-report", "MV+ report linking MTCC temporary diversion notice", "https://www.plus.mv/english/male-bus-routes-adjusted-as-boduthakurufaanu-magu-redevelopment-advances/", "Copyright MV+; local raw evidence only, article not redistributed"),
    ]
    for stem, source_id, title, url, license_name in definitions:
        metadata_file = raw/"audits"/(stem+"-metadata.json")
        if not metadata_file.exists():
            continue
        metadata = json.loads(metadata_file.read_text(encoding="utf-8"))
        file = "audits/"+stem+(".json" if source_id == "rtl-bus-routes-audit" else ".html")
        record = dict(metadata,id=source_id,title=title,url=url,license=license_name,file=file,rawPath="data/raw/"+file)
        if metadata.get("sha256") and metadata.get("status", "downloaded") == "downloaded":
            content = (raw/file).read_bytes()
            if hashlib.sha256(content).hexdigest() != metadata["sha256"] or len(content) != metadata["bytes"]:
                raise ValueError(f"Stop audit source checksum mismatch: {source_id}")
            record.update(status="downloaded",lastValidatedAt=metadata["fetchedAt"],referencePeriod=metadata["fetchedAt"][:10] if source_id == "rtl-bus-routes-audit" else "2026-09-23")
            if source_id == "rtl-bus-routes-audit":
                current_data = json.loads(content)
                if not isinstance(current_data.get("routeResponse"),list):
                    raise ValueError("Stop audit operator response schema unavailable")
            else:
                record["qualification"] = "Secondary report quoting/linking operator notice. Linked primary post could not be read; used as an unresolved temporary-diversion caution, never a verified closure."
        audit_sources.append(record)
    current_record = next((s for s in audit_sources if s["id"] == "rtl-bus-routes-audit" and s["status"] == "downloaded"),None)
    checked_at = current_record["fetchedAt"] if current_record else transit.get("fetchedAt")
    cached_routes = json.loads((raw/"rtl-bus-routes.json").read_text(encoding="utf-8")).get("routeResponse",[]) if (raw/"rtl-bus-routes.json").exists() else []
    current_routes = current_data["routeResponse"] if current_data else []
    def profile(routes):
        by_code = defaultdict(list)
        for route in routes:
            for stop in route.get("busRouteStopList",[]):
                by_code[str(stop["code"])].append(dict(stop,routeId="rtl:route:"+str(route["code"]),routeNumber=route["routeNumber"]))
        return by_code
    cached_by_code, current_by_code = profile(cached_routes), profile(current_routes)
    route_refs, pattern_refs = defaultdict(list), defaultdict(list)
    for route in transit["routes"]:
        for stop_id in set(route["stopIds"]):
            route_refs[stop_id].append(route["id"])
    for pattern in transit["patterns"]:
        for stop_id in set(pattern["stopIds"]):
            pattern_refs[stop_id].append(pattern["id"])
    locations = defaultdict(list)
    names = defaultdict(list)
    for stop in transit["stops"]:
        locations[(stop["lon"],stop["lat"])].append(stop["id"])
        names[norm(stop["name"])].append(stop["id"])
    lines = [LineString(edge["geometry"]) for edge in network["edges"] if edge.get("geometry")]
    node_coords = [node["coordinate"] for node in network["nodes"]]
    reviewed = []
    for stop in transit["stops"]:
        code = str(stop["operatorCode"])
        records = current_by_code.get(code, []) if current_record else cached_by_code.get(code, [])
        variants = {(float(row["longitude"]),float(row["latitude"])) for row in records}
        name_variants = sorted({row["name"].strip() for row in records})
        coordinate = [stop["lon"],stop["lat"]]
        point = Point(coordinate)
        coordinate_group = sorted(locations[tuple(coordinate)])
        road_offset = min((line.distance(point)*111195 for line in lines),default=None)
        node_offset = min((distance(coordinate,node) for node in node_coords),default=None)
        # No proximity connector is added. These are diagnostics against the
        # same source coastline and geometry retained by the journey engine.
        land_offset = point.distance(geoms[stop["islandId"]])*111195
        listed = code in current_by_code if current_record else None
        max_change = max((distance(coordinate,c) for c in variants),default=None)
        issues = []
        if listed is False:
            issues.append("not-in-current-operator-catalogue")
        if len(variants)>1:
            issues.append("operator-coordinate-conflict")
        if len({norm(name) for name in name_variants})>1:
            issues.append("operator-name-conflict")
        if max_change is not None and max_change>15:
            issues.append("operator-coordinate-changed-over-15m")
        if land_offset>0:
            issues.append("outside-mapped-land")
        if node_offset is None or node_offset>100:
            issues.append("road-node-attachment-unavailable")
        if not route_refs[stop["id"]]:
            issues.append("orphan-no-route-reference")
        reviewed.append(dict(id=stop["id"],operatorCode=code,name=stop["name"],islandId=stop["islandId"],coordinate=coordinate,currentOperatorListed=listed,routeIds=sorted(route_refs[stop["id"]]),patternIds=sorted(pattern_refs[stop["id"]]),timetableStatus="remaining-snapshot" if pattern_refs[stop["id"]] else "unavailable",coordinateGroupId="|".join(coordinate_group),coLocatedStopIds=coordinate_group,rawRouteStopIds=sorted({row["id"] for row in records}),operatorNameVariants=name_variants,operatorCoordinateVariantCount=len(variants),currentCoordinateDifferenceM=round(max_change,3) if max_change is not None else None,locationQA=dict(outsideLandDistanceM=round(land_offset,3),nearestRoadGeometryDistanceM=round(road_offset,3) if road_offset is not None else None,nearestGraphNodeDistanceM=round(node_offset,3) if node_offset is not None else None),physicalOperatingStatus="unverified",issues=issues,sourceIds=["rtl-bus-routes","osm-geofabrik"]+(["rtl-bus-routes-audit"] if current_record else [])))
        if "dropoffonly" in norm(stop["name"]):
            reviewed[-1].update(pickupAllowed=False,boardingRestriction="Operator stop name explicitly says DROP-OFF ONLY; alighting remains allowed.")
    loaded = sum(bool(row["patternIds"]) for row in reviewed)
    coordinate_groups = [dict(id="|".join(sorted(ids)),stopIds=sorted(ids),coordinate=list(coord),classification="published-opposite-direction-codes" if any("opp" in next(s["name"] for s in transit["stops"] if s["id"] == stop_id).lower() for stop_id in ids) else "co-located-published-codes",qualification="Preserve boarding/direction identities. Published coordinates are identical; distinct kerb positions are not independently surveyed.") for coord,ids in locations.items() if len(ids)>1]
    name_groups = [dict(normalizedName=key,stopIds=sorted(ids),qualification="Shared normalized landmark name does not prove a duplicate boarding point; inspect coordinates and route memberships.") for key,ids in names.items() if len(ids)>1]
    near_pairs = []
    for index,a in enumerate(reviewed):
        for b in reviewed[index+1:]:
            separation = distance(a["coordinate"],b["coordinate"])
            if separation<60:
                near_pairs.append(dict(stopIds=[a["id"],b["id"]],distanceM=round(separation,3),qualification="Proximity diagnostic, never an automatic merge or closure rule."))
    rows = [row for route in cached_routes for row in route["busRouteStopList"]]
    summary = dict(rawRouteStopRecords=len(rows),uniqueRouteStopRecordIds=len({row["id"] for row in rows}),uniqueOperatorStopCodes=len(cached_by_code),baselineStopIds=len(reviewed),uniqueCoordinateSites=len(locations),loadedTimingStopIds=loaded,noLoadedTimingStopIds=len(reviewed)-loaded,currentOperatorListedStopIds=sum(row["currentOperatorListed"] is True for row in reviewed),currentOperatorRemovedStopIds=sum(row["currentOperatorListed"] is False for row in reviewed),orphanStopIds=sum(not row["routeIds"] for row in reviewed),stopIdsWithFailedChecks=sum(bool(row["issues"]) for row in reviewed),dropOffOnlyStopIds=sum(row.get("pickupAllowed") is False for row in reviewed),exactCoordinateGroups=sum(len(v)>1 for v in locations.values()),duplicateNormalizedNameGroups=sum(len(v)>1 for v in names.values()),routeCatalogueCount=len(transit["routes"]),loadedPatternCount=len(transit["patterns"]),physicalClosureVerifiedStopIds=0,locationChecksQualification="Coastline/graph consistency only; all physical stops and accessibility need operator or field confirmation.")
    summary.update(currentRawRouteStopRecords=sum(len(route["busRouteStopList"]) for route in current_routes) if current_record else None,currentUniqueOperatorStopCodes=len(current_by_code) if current_record else None,newCurrentOperatorStopCodes=sorted(set(current_by_code)-set(cached_by_code)) if current_record else [],removedCurrentOperatorStopCodes=sorted(set(cached_by_code)-set(current_by_code)) if current_record else [])
    cautions = []
    report = next((s for s in audit_sources if s["id"] == "mtcc-westpark-notice-report" and s.get("status") == "downloaded"),None)
    if report:
        cautions.append(dict(id="westpark-september2026",reportedDate="2026-09-23",reviewedAt=checked_at,status="reported-temporary-diversion-unresolved",routeNumbers=["R10","R15"],primaryNoticeUrl="https://x.com/MTCCPlc/status/2102457648574812644",primaryNoticeAccess="403/inaccessible during review",sourceIds=[report["id"]],note="Report says Orchid and Orange Hiyaa temporarily bypass West Park and start/end near Maafannu Madharusa, with return diversion via Social Centre. Current route catalogue already begins those directions at Maafannu Madharusa; their timings remain unavailable. Reopening date and current driving path not verified. No stop is marked permanently closed or silently removed."))
    result = dict(schemaVersion=1,reviewedAt=checked_at,catalogueComparisonAvailable=bool(current_record),catalogueComparisonSourceId="rtl-bus-routes-audit" if current_record else None,departureSnapshotFetchedAt=transit.get("fetchedAt"),departureSnapshotDate=transit.get("snapshotDate"),summary=summary,coordinateGroups=coordinate_groups,normalizedNameGroups=name_groups,nearbyStopPairs=near_pairs,byIsland=[dict(islandId=i,operatorStopIds=sum(row["islandId"]==i for row in reviewed),loadedTimingStopIds=sum(row["islandId"]==i and bool(row["patternIds"]) for row in reviewed),noLoadedTimingStopIds=sum(row["islandId"]==i and not row["patternIds"] for row in reviewed)) for i in geoms],stops=reviewed,operationalCautions=cautions,sources=audit_sources,qualification="Unique operator stop codes preserve opposite directions and boarding identities. Route-stop IDs are repeated route memberships, not additional physical stops. Coordinate co-location is grouped visually only. Operator catalogue and remaining timings do not establish physical active/closed status or full-day service.")
    result["version"]="rtl-stop-audit-"+hashlib.sha256(json.dumps(result,sort_keys=True,ensure_ascii=False,separators=(",",":"),allow_nan=False).encode()).hexdigest()[:12]
    return result


def checked_review_sources(root, raw, records):
    """Normalize local evidence archive paths and enforce their actual bytes."""
    checked = []
    for record in records:
        source = dict(record)
        file = source.get("file") or source.get("rawPath")
        path = (root/file if file.replace("\\","/").startswith("data/raw/") else raw/file).resolve()
        if not path.is_relative_to(raw.resolve()):
            raise ValueError(f"Review evidence path outside raw archive: {source['id']}")
        if source.get("status") == "downloaded":
            content = path.read_bytes()
            if hashlib.sha256(content).hexdigest() != source["sha256"] or len(content) != source["bytes"]:
                raise ValueError(f"Review evidence changed after audit: {source['id']}; review the replacement before publishing")
        source.update(file=path.relative_to(raw.resolve()).as_posix(),rawPath="data/raw/"+path.relative_to(raw.resolve()).as_posix())
        checked.append(source)
    return checked


def apply_healthcare_review(root, raw, facilities, sources):
    file = root/"data"/"public"/"healthcare-location-audit.json"
    if not file.exists():
        raise ValueError("Healthcare location review is required; refusing to publish unaudited operating-status destinations")
    content = file.read_bytes()
    review = json.loads(content)
    if review.get("schemaVersion") != 1 or not review.get("version"):
        raise ValueError("Healthcare review schema/version unavailable")
    audit_sources = checked_review_sources(root,raw,review["sources"])
    existing = {source["id"]:source for source in sources}
    for source in audit_sources:
        if source["id"] in existing and existing[source["id"]].get("sha256") != source.get("sha256"):
            raise ValueError(f"Conflicting review evidence source ID: {source['id']}")
        if source["id"] not in existing:
            sources.append(source)
            existing[source["id"]] = source
    rows = {row["facilityId"]:row for row in review["facilities"]}
    health = [facility for facility in facilities if facility["kind"] in {"hospital","clinic"}]
    if len(rows) != len(review["facilities"]) or set(rows) != {facility["id"] for facility in health}:
        raise ValueError("Healthcare review must match every mapped healthcare ID exactly; new or missing locations require review")
    for facility in health:
        row = rows[facility["id"]]
        if row["operatingStatus"] not in {"open","closed","relocated","unverified"} or row["careScope"] not in {"general","specialist","restricted","unverified"}:
            raise ValueError(f"Healthcare status/scope invalid: {facility['id']}")
        datetime.strptime(row["statusReviewedAt"],"%Y-%m-%d")
        if not row["statusSourceIds"] or any(source_id not in existing or existing[source_id]["status"] != "downloaded" for source_id in row["statusSourceIds"]):
            raise ValueError(f"Healthcare review has no checked source evidence: {facility['id']}")
        for key in ["operatingStatus","statusReviewedAt","statusSourceIds","registryId","statusNote","careScope","careScopeNote","coordinateReview","registryEvidence"]:
            if row.get(key) is not None:
                facility[key] = row[key]
        if row.get("recommendedKind"):
            if row["recommendedKind"] not in {"hospital","clinic"}:
                raise ValueError(f"Unsupported reviewed institution category: {facility['id']}")
            facility.update(kind=row["recommendedKind"],category=row["recommendedKind"])
        facility["sourceIds"] = sorted(set(facility["sourceIds"]+row["statusSourceIds"]+row.get("reviewSourceIds",[])))
        facility["qualification"] = row["statusNote"]+" Coordinates are retained from OSM; no surveyed entrance, service capacity or step-free access certification is supplied."
    digest = hashlib.sha256(content).hexdigest()
    return dict(version=review["version"],sha256=digest,reviewedAt=review["reviewedAt"],mappedHealthcareCount=len(health),operatingStatusCounts=dict(Counter(row["operatingStatus"] for row in rows.values())),generalCareEligibleCount=sum(row["operatingStatus"] == "open" and row["careScope"] == "general" and next(f for f in health if f["id"]==row["facilityId"])["verification"] is not False for row in rows.values()),coverageGaps=review.get("coverageGaps",[]),limitations=review.get("limitations",[]))


def apply_facility_operation_review(root, raw, facilities, sources):
    """Require a current evidence decision for every education location, including exclusions."""
    file = root/"data"/"public"/"facility-operation-review.json"
    if not file.exists():
        raise ValueError("Facility operation review required; refusing to publish unaudited school destinations")
    content = file.read_bytes()
    review = json.loads(content)
    if review.get("schemaVersion") != 1 or not review.get("version"):
        raise ValueError("Facility operation review schema/version unavailable")
    datetime.strptime(review["reviewedAt"],"%Y-%m-%d")
    existing = {source["id"]:source for source in sources}
    for source in checked_review_sources(root,raw,review["sources"]):
        if source["id"] in existing and existing[source["id"]].get("sha256") != source.get("sha256"):
            raise ValueError(f"Conflicting operation review evidence: {source['id']}")
        if source["id"] not in existing:
            sources.append(source)
            existing[source["id"]] = source
    rows = {row["facilityId"]:row for row in review["records"]}
    education_ids = {facility["id"] for facility in facilities if facility["kind"] in {"school","college"}}
    facility_ids = {facility["id"] for facility in facilities}
    if len(rows) != len(review["records"]) or not education_ids.issubset(rows) or not set(rows).issubset(facility_ids):
        raise ValueError("Operation review must cover every education ID; new or missing locations require review")
    for facility in facilities:
        row = rows.get(facility["id"])
        if not row:
            continue
        if row["operatingStatus"] not in {"open","closed","relocated","unverified"}:
            raise ValueError(f"Invalid reviewed operating status: {facility['id']}")
        datetime.strptime(row["statusReviewedAt"],"%Y-%m-%d")
        if not row["statusSourceIds"] or any(source_id not in existing or existing[source_id]["status"] != "downloaded" for source_id in row["statusSourceIds"]):
            raise ValueError(f"Operation review lacks checked evidence: {facility['id']}")
        for key in ["operatingStatus","statusReviewedAt","statusSourceIds","statusNote","coordinateReview","educationSuitability","institutionStatus"]:
            if row.get(key) is not None:
                facility[key] = row[key]
        # Review prose qualifies institutional evidence, not a new routing enum.
        # Retain the original mapped-location policy; explicit exclusions fail closed.
        if row.get("verification") is False:
            facility["verification"] = False
        elif row.get("verification") is not None:
            facility["verificationNote"] = row["verification"]
        if row.get("recommendedKind"):
            if row["recommendedKind"] not in {"school","college"} or facility["id"] not in education_ids:
                raise ValueError(f"Unsupported reviewed education category: {facility['id']}")
            facility.update(kind=row["recommendedKind"],category=row["recommendedKind"])
        if row.get("generalSchoolEligible") is False and facility["kind"] == "school":
            facility.update(verification=False,educationSuitability="excluded")
        facility["sourceIds"] = sorted(set(facility["sourceIds"]+row["statusSourceIds"]))
        facility["qualification"] = row["statusNote"]+" Desk review is not an inspection of entrances, operating hours, admissions or capacity."
    education = [facility for facility in facilities if facility["id"] in education_ids]
    return dict(version=review["version"],sha256=hashlib.sha256(content).hexdigest(),reviewedAt=review["reviewedAt"],mappedEducationCount=len(education),operatingStatusCounts=dict(Counter(facility["operatingStatus"] for facility in education)),generalSchoolEligibleCount=sum(facility["kind"] == "school" and facility["operatingStatus"] == "open" and facility["verification"] is not False for facility in education),sourceIds=sorted({source_id for row in review["records"] for source_id in row["statusSourceIds"]}),qualification=review.get("qualification","Positive institution and mapped-site evidence required; uncertain sites stay outside default access results."))


def apply_transit_operation_review(root, raw, transit, sources):
    """Use a fresh catalogue audit without replacing the validated earlier service snapshot."""
    file = root/"data"/"public"/"transit-operation-review.json"
    content = file.read_bytes()
    review = json.loads(content)
    if review.get("schemaVersion") != 1 or not review.get("version"):
        raise ValueError("Transit operation review schema/version unavailable")
    records = []
    for original in review["sources"]:
        if not original.get("rawPath") or not original.get("sha256"):
            continue  # Context and inaccessible notices are not downloaded evidence.
        source = dict(original)
        path = root/source["rawPath"]
        source.update(status="downloaded",title="RTL fresh operator route catalogue check",fetchedAt=source["retrievedAt"],bytes=path.stat().st_size)
        records.append(source)
    if not records:
        raise ValueError("Fresh transit review requires archived catalogue evidence")
    catalogue = json.loads((root/records[0]["rawPath"]).read_text(encoding="utf-8"))
    current_positions = defaultdict(set)
    for route in catalogue["routeResponse"]:
        for stop in route["busRouteStopList"]:
            current_positions["rtl:stop:"+str(stop["code"])].add((float(stop["longitude"]),float(stop["latitude"])))
    existing = {source["id"]:source for source in sources}
    for source in checked_review_sources(root,raw,records):
        if source["id"] in existing and existing[source["id"]].get("sha256") != source["sha256"]:
            raise ValueError(f"Conflicting transit operation source: {source['id']}")
        if source["id"] not in existing:
            sources.append(source)
    rows = {row["id"]:row for row in review["stops"]}
    if len(rows) != len(review["stops"]) or set(rows) != {stop["id"] for stop in transit["stops"]} or set(rows) != set(current_positions):
        raise ValueError("Fresh operator review must match the entire stop inventory")
    for stop in transit["stops"]:
        row = rows[stop["id"]]
        if not row["publishedCoordinateMatchesBaseline"] or current_positions[stop["id"]] != {tuple(stop["coordinate"])}:
            raise ValueError(f"Operator stop moved; review coordinates before publication: {stop['id']}")
        stop.update(currentOperatorListed=row["currentOperatorListed"],catalogueReviewedAt=review["reviewedAt"],physicalOperatingStatus=row["physicalOperatingStatus"])
        stop["sourceIds"] = sorted(set(stop["sourceIds"]+[source["id"] for source in records]))
    transit["operationReview"] = dict(version=review["version"],reviewedAt=review["reviewedAt"],summary=review["summary"],byIsland=review["byIsland"],departureCoverage=review["departureCoverage"],operationalCautions=review["operationalCautions"],qualification="Current operator catalogue listing does not establish field operation. Untimed listings remain hidden by default; original validated departure data is retained.")
    return hashlib.sha256(content).hexdigest()


def make_education_review(root, raw, facilities, osm, geoms, network, sources):
    definitions = ["education-directory-current","fiyavathi-presidency2024","stem-maaharaa2026"]
    metadata = []
    for stem in definitions:
        record = next((source for source in sources if source.get("file") == "audits/"+stem+".html" and source.get("status") == "downloaded"),None)
        if not record:
            file = raw/"audits"/(stem+"-metadata.json")
            if not file.exists():
                raise ValueError(f"Education review evidence required: {stem}; run fetch")
            record = json.loads(file.read_text(encoding="utf-8"))
        metadata.append(record)
    audit_sources = checked_review_sources(root,raw,metadata)
    ids = {source["id"] for source in sources}
    sources += [source for source in audit_sources if source["id"] not in ids]
    soup = BeautifulSoup((raw/"audits"/"education-directory-current.html").read_text(encoding="utf-8"),"html.parser")
    names = {norm(x.get_text(" ",strip=True)) for x in soup.select("span.font-bold") if len(norm(x.get_text()))>5}
    exclusions = {
        "osm:way:1210925026":dict(sourceId="education-fiyavathi-presidency",note="Primary President's Office record identifies Fiyavathi as a state-run children's care facility. General public school enrollment is not established; excluded from School access while retaining the mapped inventory location."),
        "osm:node:11183603641":dict(sourceId="education-stem-maaharaa",note="Ministry of Health registers Stem Center as a therapy service (HF-22-ALP/PVT/675, PVT-TS-22-335-00). General public school suitability is not established; excluded from School access. OSM school category retained as source inventory, not recast as general healthcare."),
    }
    tagmap = {f"osm:{poi['osmType']}:{poi['osmId']}":poi["tags"] for poi in osm.pois}
    lines = [LineString(edge["geometry"]) for edge in network["edges"] if edge.get("geometry")]
    records = []
    for facility in facilities:
        if facility["kind"] not in {"school","college"}:
            continue
        exclusion = exclusions.get(facility["id"])
        if exclusion:
            facility.update(verification=False,statusNote=exclusion["note"],qualification=exclusion["note"],educationSuitability="excluded")
            facility["sourceIds"] = sorted(set(facility["sourceIds"]+[exclusion["sourceId"]]))
        tags = tagmap[facility["id"]]
        point = Point(facility["coordinate"])
        records.append(dict(id=facility["id"],name=facility["name"],kind=facility["kind"],islandId=facility["islandId"],coordinate=facility["coordinate"],officialDirectoryNameMatched=norm(facility["name"]) in names,operatingStatus="not-audited",generalSchoolEligible=facility["kind"] == "school" and facility["verification"] is not False,excluded=bool(exclusion),exclusionReason=exclusion["note"] if exclusion else None,sourceIds=facility["sourceIds"]+["education-directory-audit"],coordinateQA=dict(outsideLandDistanceM=round(point.distance(geoms[facility["islandId"]])*111195,3),nearestRoadGeometryDistanceM=round(min(line.distance(point) for line in lines)*111195,3),nearestGraphNodeDistanceM=round(min(distance(facility["coordinate"],node["coordinate"]) for node in network["nodes"]),3)),osmLifecycleTags={key:value for key,value in tags.items() if any(term in key for term in ["disused","abandon","closed","start_date","end_date","check_date","opening_hours","construction","access"])}))
    summary = dict(mappedEducationPoints=len(records),mappedSchoolPoints=sum(record["kind"] == "school" for record in records),mappedCollegePoints=sum(record["kind"] == "college" for record in records),officialDirectoryNameMatches=sum(record["officialDirectoryNameMatched"] for record in records),osmOnlyPoints=sum(not record["officialDirectoryNameMatched"] for record in records),generalSchoolEligiblePoints=sum(record["generalSchoolEligible"] for record in records),excludedSchoolPoints=sum(record["excluded"] for record in records),outsideMappedLandPoints=sum(record["coordinateQA"]["outsideLandDistanceM"]>0 for record in records),roadAttachmentUnavailablePoints=sum(record["coordinateQA"]["nearestGraphNodeDistanceM"]>100 for record in records),duplicateNormalizedNameGroups=sum(count>1 for count in Counter(norm(record["name"]) for record in records).values()))
    result = dict(schemaVersion=1,reviewedAt=max(source["fetchedAt"] for source in audit_sources),directoryPublishedUpdate="2025-05-23",summary=summary,excludedIds=sorted(exclusions),sourceIds=[source["id"] for source in audit_sources],sources=audit_sources,records=records,qualification="MoE government-school directory names checked against all mapped education points. Page displays update23 May2025; absence does not establish private-school/college closure. Coordinates are OSM-derived with coastline/graph consistency diagnostics, not independently surveyed entrances. Colleges are inventory only for School goal. Current operating status/capacity/admission for the other mapped schools has not been audited.")
    result["version"]="education-location-audit-"+hashlib.sha256(json.dumps(result,sort_keys=True,ensure_ascii=False,separators=(",",":"),allow_nan=False).encode()).hexdigest()[:12]
    return result


def prepare_snapshot(root, raw, public, manifest_path, output, bounds):
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    sources = manifest["sources"]
    for source in sources:
        if source.get("required") and source["status"] != "downloaded":
            raise ValueError(f"Required source unavailable: {source['id']}")
        if source["status"] == "downloaded":
            actual = hashlib.sha256((raw/source["file"]).read_bytes()).hexdigest()
            if actual != source["sha256"]:
                raise ValueError(f"Raw source checksum mismatch: {source['id']}; run fetch --refresh")
        source["referencePeriod"] = source.get("referenceYear")
    osm = OsmSnapshot(bounds)
    osm.apply_file(str(raw/"maldives-latest.osm.pbf"), locations=True)
    islands, geoms = build_islands(osm, raw)
    demographics = make_demographics(islands,raw)
    origins, origin_report = make_origins(islands, geoms, raw)
    network = make_network(osm, geoms)
    available = {s["id"] for s in sources if s["status"] == "downloaded"}
    facilities, facility_exclusions = make_facilities(osm, geoms, raw, available)
    healthcare_review = apply_healthcare_review(root, raw, facilities, sources)
    education_review = make_education_review(root, raw, facilities, osm, geoms, network, sources)
    operation_review = apply_facility_operation_review(root, raw, facilities, sources)
    for record in education_review["records"]:
        facility = next(f for f in facilities if f["id"] == record["id"])
        record.update(kind=facility["kind"],operatingStatus=facility["operatingStatus"],statusReviewedAt=facility["statusReviewedAt"],statusSourceIds=facility["statusSourceIds"],statusNote=facility["statusNote"],generalSchoolEligible=facility["kind"] == "school" and facility["operatingStatus"] == "open" and facility["verification"] is not False)
        record.update(excluded=facility["kind"] == "school" and not record["generalSchoolEligible"],sourceIds=facility["sourceIds"])
        record["exclusionReason"] = facility["statusNote"] if record["excluded"] else None
    education_review["summary"]["generalSchoolEligiblePoints"] = operation_review["generalSchoolEligibleCount"]
    education_review["excludedIds"] = sorted(record["id"] for record in education_review["records"] if record["excluded"])
    education_review["summary"]["excludedSchoolPoints"] = len(education_review["excludedIds"])
    education_review["sourceIds"] = sorted(set(education_review["sourceIds"]+operation_review["sourceIds"]))
    education_review["sources"] = [source for source in sources if source["id"] in education_review["sourceIds"]]
    education_review["operationReviewVersion"] = operation_review["version"]
    education_review["qualification"] = "OSM locations reconciled with institution evidence and operating-status review. Relocated, closed and unverified mapped sites are excluded from School access. Coordinates are not independently surveyed entrances; colleges are inventory only."
    education_review["version"] = "education-location-audit-"+hashlib.sha256(json.dumps(education_review,sort_keys=True,ensure_ascii=False,separators=(",",":"),allow_nan=False).encode()).hexdigest()[:12]
    transit, reference_date = make_transit(raw, geoms, sources, network)
    stop_audit = audit_transit_stops(transit, network, geoms, raw, sources)
    source_ids = {source["id"] for source in sources}
    sources += [source for source in stop_audit["sources"] if source["id"] not in source_ids]
    stop_reviews = {review["id"]:review for review in stop_audit["stops"]}
    for stop in transit["stops"]:
        review = stop_reviews[stop["id"]]
        stop.update({key:review[key] for key in ["routeIds","patternIds","timetableStatus","coordinateGroupId","currentOperatorListed"]})
        stop["catalogueReviewedAt"] = stop_audit["reviewedAt"]
        if review["currentOperatorListed"] is not None:
            stop["sourceIds"].append("rtl-bus-routes-audit")
    transit["stopAudit"] = dict(version=stop_audit["version"],reviewedAt=stop_audit["reviewedAt"],summary=stop_audit["summary"],coordinateGroups=stop_audit["coordinateGroups"],operationalCautions=stop_audit["operationalCautions"],qualification=stop_audit["qualification"])
    transit_operation_sha = apply_transit_operation_review(root, raw, transit, sources)
    boundaries = dict(type="FeatureCollection", features=[dict(type="Feature", id=i["id"], properties={k:v for k,v in i.items() if k not in {"geometry", "bounds"}}, geometry=i["geometry"]) for i in islands])
    buildings = dict(type="FeatureCollection", features=[])
    for b in osm.buildings:
        island_id = locate_island((b["geometry"].representative_point().x,b["geometry"].representative_point().y), geoms)
        if island_id:
            buildings["features"].append(dict(type="Feature", id="osm:way:"+str(b["id"]), properties=dict(islandId=island_id, levels=b["tags"].get("building:levels"), name=b["tags"].get("name:en") or b["tags"].get("name"), sourceId="osm-geofabrik"), geometry=mapping(b["geometry"].simplify(0.000002,preserve_topology=True))))
    preparation_sha = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    acquisition_sha = hashlib.sha256((root/"scripts"/"public-data.py").read_bytes()).hexdigest()
    snapshot_utc = max(s["fetchedAt"] for s in sources if s["status"] == "downloaded")
    snapshot_fingerprint = "".join(s["id"]+s.get("sha256","")+s.get("fetchedAt","") for s in sources)+preparation_sha+acquisition_sha+stop_audit["version"]+healthcare_review["sha256"]+education_review["version"]+operation_review["sha256"]+transit_operation_sha
    version = "greater-male-"+(reference_date or datetime.now(timezone.utc).date().isoformat())+"-"+hashlib.sha256(snapshot_fingerprint.encode()).hexdigest()[:10]
    cost_anchors = make_cost_anchors(raw,available)
    baseline = dict(schemaVersion=1, version=version, generatedAt=now(), referenceDate=reference_date or datetime.now(timezone.utc).date().isoformat(), studyArea=dict(name="Greater Malé",bounds=list(bounds),populationScope="Malé four wards, Hulhumalé, Vilimalé. Hulhulé mapped but projected resident allocation unavailable; harbours and combined outer-industrial-island projection excluded."), sources=sources, islands=islands, islandBoundaries=boundaries, landBoundaries=boundaries, origins=origins, network=network, facilities=facilities, destinations=facilities, transit=transit, buildings=buildings, costAnchors=cost_anchors, assumptions=dict(populationSpatialDistribution="WorldPop 2026 alpha, land-masked and normalized to MBS projections; 2030 shares 2026 distribution",hulhumaleAirportPartitionLatitude=4.203,turnRestrictions="Not included",motorOneway="Preserved",sourceGaps=["No observed travel demand, ridership, route-specific operating cost or mode-shift elasticity supplied", "Official age cohorts from administrative island indicator workbook do not cover Malé; no invented neighborhood demographics", "Hulhulé 2026 and 2030 are grouped with Thilafushi/Gulhifalhu in official projections and cannot be independently allocated"]), validation=dict(originNormalization=origin_report,facilityExclusions=facility_exclusions,lastValidatedAt=now()))
    baseline["sourceSnapshotUtc"] = snapshot_utc
    baseline["methodology"] = dict(version="greater-male-1.1",preparationSha256=preparation_sha,acquisitionSha256=acquisition_sha)
    baseline["demographics"] = demographics
    baseline["facilityAuditVersion"] = healthcare_review["version"]
    baseline["healthcareLocationAudit"] = healthcare_review
    baseline["facilityOperationReview"] = operation_review
    baseline["educationLocationAudit"] = {key:education_review[key] for key in ["version","reviewedAt","directoryPublishedUpdate","summary","excludedIds","sourceIds","qualification"]}
    baseline["methodology"].update(healthcareAuditSha256=healthcare_review["sha256"],educationAuditVersion=education_review["version"],facilityOperationReviewSha256=operation_review["sha256"],transitAuditVersion=stop_audit["version"],transitOperationReviewSha256=transit_operation_sha)
    candidate = output.with_name(output.stem+".candidate"+output.suffix)
    try:
        write_json(candidate, baseline)
        validate_snapshot(candidate)
        candidate.replace(output)
    finally:
        candidate.unlink(missing_ok=True)
    write_json(public/"greater-male-boundaries.geojson", boundaries)
    write_json(public/"greater-male-buildings.geojson", buildings)
    write_json(public/"source-manifest.json", dict(version=version,generatedAt=now(),sourceSnapshotUtc=snapshot_utc,methodology=baseline["methodology"],serviceCoverage=transit.get("serviceCoverage"),sources=sources))
    write_json(public/"transit-stop-audit.json", stop_audit)
    write_json(public/"education-location-audit.json", education_review)
    print(json.dumps(dict(output=str(output),version=version,islands=len(islands),origins=len(origins),nodes=len(network["nodes"]),edges=len(network["edges"]),facilities=len(facilities),buildings=len(buildings["features"]),stops=len(transit["stops"]),patterns=len(transit["patterns"]),bytes=output.stat().st_size)))


def validate_snapshot(path):
    d = json.loads(path.read_text(encoding="utf-8"))
    assert d["schemaVersion"] == 1
    islands = {i["id"]:i for i in d["islands"]}
    geoms = {k:shape(v["geometry"]) for k,v in islands.items()}
    nodes = {n["id"]:n for n in d["network"]["nodes"]}
    sources = {s["id"]:s for s in d["sources"]}
    assert len(nodes) == len(d["network"]["nodes"])
    assert all(e["from"] in nodes and e["to"] in nodes and e["distanceM"]>0 for e in d["network"]["edges"])
    assert all(e["distanceM"] <= 70.01 for e in d["network"]["edges"]), "Road attachment spacing exceeds source-geometry split limit"
    assert len({e["id"] for e in d["network"]["edges"]}) == len(d["network"]["edges"])
    for record in d["islands"]+d["origins"]+d["facilities"]+d["transit"]["stops"]+d["transit"]["patterns"]:
        assert all(source_id in sources and sources[source_id]["status"] == "downloaded" for source_id in record["sourceIds"]), f"Unsourced record {record['id']}"
    for facility in d["facilities"]:
        if facility["verification"] == "official":
            assert any(source_id != "osm-geofabrik" for source_id in facility["sourceIds"]), f"Official match without directory source: {facility['name']}"
        if facility["kind"] in {"hospital","clinic"}:
            assert d.get("facilityAuditVersion"), "Audited public healthcare requires review version"
            assert facility.get("operatingStatus") in {"open","closed","relocated","unverified"}
            assert facility.get("careScope") in {"general","specialist","restricted","unverified"}
            datetime.strptime(facility["statusReviewedAt"],"%Y-%m-%d")
            assert facility["statusSourceIds"] and all(source_id in sources and sources[source_id]["status"] == "downloaded" for source_id in facility["statusSourceIds"])
        if facility["kind"] in {"school","college"}:
            assert d.get("facilityOperationReview"), "Public education requires an operation review"
            assert facility.get("operatingStatus") in {"open","closed","relocated","unverified"}
            datetime.strptime(facility["statusReviewedAt"],"%Y-%m-%d")
            assert facility["statusSourceIds"] and all(source_id in sources and sources[source_id]["status"] == "downloaded" for source_id in facility["statusSourceIds"])
    assert d["methodology"].get("healthcareAuditSha256"), "Healthcare review checksum must participate in snapshot version"
    assert d["methodology"].get("facilityOperationReviewSha256"), "Education operation review must participate in snapshot version"
    for record_id in d["educationLocationAudit"]["excludedIds"]:
        facility = next(f for f in d["facilities"] if f["id"] == record_id)
        assert facility["verification"] is False and facility.get("statusNote"), "Reviewed non-school inventory point must be excluded from School access"
    for o in d["origins"]:
        assert o["population2026"]>0 and o["population2030"]>0
        # Minor simplification and seven-decimal point rounding can differ by
        # centimetres; the actual clipped origin footprint remains on source land.
        assert geoms[o["islandId"]].buffer(0.00001).covers(Point(o["lon"],o["lat"])), f"Offshore origin {o['id']}"
    for island_id,island in islands.items():
        for year in [2026,2030]:
            if island[f"population{year}"] is not None:
                total=sum(o[f"population{year}"] for o in d["origins"] if o["islandId"]==island_id)
                assert abs(total-island[f"population{year}"]) < 0.001, f"Population mismatch {island_id}/{year}"
    for island in d["demographics"]["islands"]:
        for group in island["availableGroups"]:
            assert not group["canAllocateToOrigins"]
            assert abs(sum(g["value"] for g in group["groups"])-group["denominator"]) < 0.01, f"Cohort total mismatch {group['id']}"
            assert all(g["value"] >= 0 for g in group["groups"])
            assert all(s in sources and sources[s]["status"] == "downloaded" for s in group["sourceIds"])
    stops = {s["id"] for s in d["transit"]["stops"]}
    assert len(stops) == len(d["transit"]["stops"]), "Duplicate canonical operator stop ID"
    if d["transit"].get("stopAudit"):
        summary = d["transit"]["stopAudit"]["summary"]
        timed = {stop_id for pattern in d["transit"]["patterns"] for stop_id in pattern["stopIds"]}
        assert summary["baselineStopIds"] == len(stops)
        assert summary["loadedTimingStopIds"] == len(timed)
        assert summary["noLoadedTimingStopIds"] == len(stops-timed)
        assert summary["dropOffOnlyStopIds"] == sum(s.get("pickupAllowed") is False for s in d["transit"]["stops"])
        for stop in d["transit"]["stops"]:
            if "dropoffonly" in norm(stop["name"]):
                assert stop.get("pickupAllowed") is False, "Published drop-off-only stop must not allow boarding"
            assert bool(stop["patternIds"]) == (stop["id"] in timed)
    for p in d["transit"]["patterns"]:
        assert set(p["stopIds"]).issubset(stops)
        assert p["sourceIds"] and p["provenance"] == "observed"
        assert p["calendar"]["includedDates"] == [d["transit"]["snapshotDate"]] and p["calendar"]["weekdays"] == [], "Remaining-departure API must not become recurring timetable"
        assert len(p["segmentGeometry"]) == len(p["stopIds"])-1
        assert len(p["segmentDistanceM"]) == len(p["stopIds"])-1
        if p["distanceM"] is not None:
            assert abs(p["distanceM"]-sum(p["segmentDistanceM"])) < 0.01
        for trip in p["timing"].get("trips",[]):
            assert len(trip["stopTimes"]) == len(p["stopIds"])
            times=[t["departureMinutes"] for t in trip["stopTimes"]]
            assert all(b>=a for a,b in zip(times,times[1:]))
    assert not any(o["islandId"]=="hulhule" for o in d["origins"]), "Do not invent grouped Hulhulé projection"
    assert all(s.get("sha256") and s.get("fetchedAt") for s in d["sources"] if s["status"]=="downloaded")
    print(f"VALIDATED: {len(d['origins'])} land-masked origins, conserved 2026/2030 island totals, {len(nodes)} topology-preserving road nodes, source joins and snapshot-only schedules.",flush=True)
