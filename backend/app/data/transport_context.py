"""Published monthly transport categories with missing and quarantined cells intact."""

from __future__ import annotations

import hashlib
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

URL = "https://statisticsmaldives.gov.mv/mbs/wp-content/uploads/2026/09/MIF_June-2026.xlsx"
EXPECTED_SHA256 = "e591c7a3adbf20937a7fea49ec5970fb28f2aa25adf53bea191436b26e15bfaa"
MONTHS = ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec")


def prepare(root: Path) -> dict[str, Any]:
    path = root / "data/raw/mbs-monthly-june2026.xlsx"
    content = path.read_bytes()
    digest = hashlib.sha256(content).hexdigest()
    if digest != EXPECTED_SHA256:
        raise ValueError("Monthly workbook differs from the reviewed June 2026 source vintage")
    audit_path = root / "data/public/prediction-source.json"
    audit = json.loads(audit_path.read_text(encoding="utf-8"))
    repeated = audit["audit"]["repeatedBusValues"]
    assert {r["source_row"] for r in repeated} == set(range(166, 172))
    assert all(r["equal"] and r["matches_prior_may_release"] for r in repeated)
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["June 2026"]
    rows = list(sheet.values)
    observations = []
    for source_row in (*range(150, 157), *range(159, 165), *range(166, 172)):
        source_values = rows[source_row - 1]
        label = source_values[2]
        if not isinstance(label, str) or not label.strip():
            raise ValueError(f"Missing category label in row {source_row}")
        table = "11.1" if source_row <= 156 else "11.2" if source_row <= 164 else "11.3"
        mode = "bus" if table == "11.3" else "ferry"
        for offset in range(78):
            column_index = 3 + offset
            year, month = 2020 + offset // 12, offset % 12 + 1
            header = rows[147][column_index]
            actual_month = (
                header.month
                if isinstance(header, (date, datetime))
                else (MONTHS.index(str(header).strip().casefold()[:3]) + 1)
            )
            if actual_month != month:
                raise ValueError(f"Unexpected month header at column {column_index + 1}")
            published = source_values[column_index]
            numeric = isinstance(published, (int, float)) and not isinstance(published, bool)
            if numeric and (published < 0 or int(published) != published):
                raise ValueError("Passenger movement cell must be a nonnegative count")
            status = "reported" if numeric else "blank" if published is None else "unavailable"
            value = int(published) if numeric else None
            quarantine = mode == "bus" and year == 2026 and month == 6
            if quarantine:
                if published != rows[source_row - 1][column_index - 1]:
                    raise ValueError("Reviewed June bus duplicate condition changed")
                status, value = "quarantined_repeated_month", None
            observations.append(
                {
                    "id": f"mbs-transport:{source_row}:{year}-{month:02}",
                    "category_id": f"mbs-transport-row-{source_row}",
                    "category": label.strip(),
                    "table": table,
                    "mode": mode,
                    "geography_grain": "published_atoll_zone"
                    if table == "11.1"
                    else "published_operator_category",
                    "period": f"{year}-{month:02}",
                    "year": year,
                    "month": month,
                    "value": value,
                    "published_value": published,
                    "value_status": status,
                    "unit": "passenger_movements",
                    "forecast_eligible": False,
                    "source_id": "mbs-monthly-june2026",
                    "source_sheet": sheet.title,
                    "source_row": source_row,
                    "source_cell": f"{get_column_letter(column_index + 1)}{source_row}",
                    "qualification": "June bus cells repeat May in all six categories and are "
                    "quarantined, as in the reviewed passenger source."
                    if quarantine
                    else "Published category totals may overlap. They are not unique people, "
                    "island demand or route-level boardings.",
                }
            )
    workbook.close()
    return {
        "schemaVersion": 1,
        "version": "official-transport-context-v1:" + digest[:16],
        "sourceSnapshotDate": "2026-10-16",
        "sources": [
            {
                "id": "mbs-monthly-june2026",
                "title": "MBS Maldives in Figures, June 2026, transport tables 11.1–11.3",
                "url": URL,
                "publisher": "Maldives Bureau of Statistics",
                "referencePeriod": "2020-01 through 2026-06; source blanks and NA retained",
                "sha256": digest,
                "rawPath": path.relative_to(root).as_posix(),
                "licenseStatus": "Official public download; explicit reuse licence not stated",
            }
        ],
        "sourceReceipts": [
            {
                "path": p.relative_to(root).as_posix(),
                "sha256": hashlib.sha256(p.read_bytes()).hexdigest(),
                "bytes": p.stat().st_size,
            }
            for p in (path, audit_path)
        ],
        "observations": observations,
        "coverage": {
            "categories": 19,
            "cells": len(observations),
            "reported_values": sum(r["value"] is not None for r in observations),
            "quarantined_cells": sum(
                r["value_status"].startswith("quarantined") for r in observations
            ),
        },
        "warnings": [
            "Atoll zones are the publisher's labels, not island boundaries. Two distinct "
            "Zone 4 rows retain separate source-row identities.",
            "Ferry endpoints, zone totals and free-of-charge bus categories can overlap. "
            "Do not sum categories into unique riders or population.",
            "These context categories do not extend the application's three reviewed "
            "forecast series; forecast eligibility remains separate.",
        ],
    }
