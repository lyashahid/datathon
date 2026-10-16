"""Census constituency context with source units and unmatched boundaries preserved.

Source-row keys distinguish two identically named Meedhoo constituencies. They are
not geographic IDs and must never be used to allocate values to island origins.
"""

from __future__ import annotations

import hashlib
import json
import math
from collections import Counter
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

CONTEXT_VERSION = "official-constituency-context-v1"
KINDS = ("population", "employment", "education")
FIRST_ROW = 5
LAST_ROW = 97
YEAR = 2022
GRAIN_WARNING = (
    "Constituency boundaries are unmatched. Source-row keys are not island IDs; "
    "do not allocate these values to islands, resident origins or travel demand."
)


@dataclass(frozen=True)
class Metric:
    column: str
    metric: str
    population_scope: str = "resident Maldivians"
    sex: str = "all"
    age_group: str = "all"
    unit: str = "people"
    published_unit: str = "people"
    denominator: str = "not applicable"


POPULATION_METRICS = (
    Metric("B", "resident_population", "all residents"),
    Metric("C", "resident_population_male", "all residents", "male"),
    Metric("D", "resident_population_female", "all residents", "female"),
    Metric("E", "maldivian_population"),
    Metric("F", "maldivian_population_male", sex="male"),
    Metric("G", "maldivian_population_female", sex="female"),
    Metric("H", "foreign_population", "resident foreigners"),
    Metric("I", "foreign_population_male", "resident foreigners", "male"),
    Metric("J", "foreign_population_female", "resident foreigners", "female"),
    Metric(
        "K",
        "resident_foreign_percent",
        "all residents",
        unit="percent",
        published_unit="percent",
        denominator="all resident population",
    ),
    Metric(
        "L",
        "maldivian_sex_ratio",
        unit="Maldivian males per 100 females",
        published_unit="Maldivian males per 100 females",
        denominator="resident Maldivian females",
    ),
    Metric("M", "maldivian_children", age_group="0-17"),
    Metric("N", "maldivian_youth", age_group="18-35"),
    Metric("O", "maldivian_working_age_36_64", age_group="36-64"),
    Metric("P", "maldivian_elderly", age_group="65+"),
    *(
        Metric(
            column,
            metric,
            age_group=age,
            unit="percent",
            published_unit="ratio",
            denominator="all resident Maldivians",
        )
        for column, metric, age in (
            ("Q", "maldivian_children_percent", "0-17"),
            ("R", "maldivian_youth_percent", "18-35"),
            ("S", "maldivian_working_age_36_64_percent", "36-64"),
            ("T", "maldivian_elderly_percent", "65+"),
        )
    ),
)

EMPLOYMENT_METRICS = (
    Metric("B", "employed", age_group="15+"),
    Metric("C", "unemployed", age_group="15+"),
    Metric("D", "outside_labour_force", age_group="15+"),
    Metric("E", "labour_status_population", age_group="15+"),
    *(
        Metric(column, metric, sex=sex, age_group=age)
        for column, metric, sex, age in (
            ("F", "youth_15_24_male", "male", "15-24"),
            ("G", "youth_15_24_female", "female", "15-24"),
            ("H", "youth_15_24_total", "all", "15-24"),
            ("I", "neet_15_24_male", "male", "15-24"),
            ("J", "neet_15_24_female", "female", "15-24"),
            ("K", "neet_15_24_total", "all", "15-24"),
            ("L", "youth_18_35_male", "male", "18-35"),
            ("M", "youth_18_35_female", "female", "18-35"),
            ("N", "youth_18_35_total", "all", "18-35"),
            ("O", "neet_18_35_male", "male", "18-35"),
            ("P", "neet_18_35_female", "female", "18-35"),
            ("Q", "neet_18_35_total", "all", "18-35"),
        )
    ),
    Metric(
        "R",
        "unemployment_percent",
        age_group="15+",
        unit="percent",
        published_unit="ratio",
        denominator="employed plus unemployed Maldivians aged 15+",
    ),
    Metric(
        "S",
        "neet_15_24_percent",
        age_group="15-24",
        unit="percent",
        published_unit="ratio",
        denominator="resident Maldivians aged 15-24",
    ),
    Metric(
        "T",
        "neet_18_35_percent",
        age_group="18-35",
        unit="percent",
        published_unit="ratio",
        denominator="resident Maldivians aged 18-35",
    ),
)

ATTAINMENTS = (
    ("not_stated", "Not stated"),
    ("preschool", "Pre-school education"),
    ("primary", "Primary education"),
    ("secondary", "Secondary education"),
    ("higher_secondary", "Higher secondary education"),
    ("certificate_diploma", "Certificate/ Diploma"),
    ("degree_and_above", "Degree & above"),
    ("level_not_stated", "Level not stated"),
    ("never_attended", "Never attended"),
)
EDUCATION_METRICS = (
    Metric("B", "literacy_population_10_plus", age_group="10+"),
    Metric("C", "mother_tongue_literate", age_group="10+"),
    Metric(
        "D",
        "mother_tongue_literacy_percent",
        age_group="10+",
        unit="percent",
        published_unit="ratio",
        denominator="resident Maldivians aged 10+",
    ),
    Metric("E", "english_literate", age_group="10+"),
    Metric(
        "F",
        "english_literacy_percent",
        age_group="10+",
        unit="percent",
        published_unit="ratio",
        denominator="resident Maldivians aged 10+",
    ),
    Metric("G", "attainment_population_15_plus", age_group="15+"),
    *(
        Metric(get_column_letter(column), f"attainment_{metric}", age_group="15+")
        for column, (metric, _) in enumerate(ATTAINMENTS, 8)
    ),
    *(
        Metric(
            get_column_letter(column),
            f"attainment_{metric}_percent",
            age_group="15+",
            unit="percent",
            published_unit="ratio",
            denominator="resident Maldivians aged 15+",
        )
        for column, (metric, _) in enumerate(ATTAINMENTS[1:], 17)
    ),
)

METRICS = {
    "population": POPULATION_METRICS,
    "employment": EMPLOYMENT_METRICS,
    "education": EDUCATION_METRICS,
}
TITLES = {
    "population": "Population counted in the constituency, Census 2022",
    "employment": "Maldivian Population counted in the constituency by employment status, "
    "Census 2022",
    "education": "Maldivian Population counted in the constituency by education status, "
    "Census 2022",
}
GROUP_HEADERS = {
    "population": {
        "A": "Constituency Area",
        "B": "Resident  Population",
        "E": "Resident Maldivian Population",
        "H": "Resident Foreign Population",
        "K": "Resident Foreigners (%) (out of resident Population",
        "L": "Sex Ratio",
        "M": "Maldivian Population by broad age ",
        "Q": "Maldivian Population by broad age (%)",
    },
    "employment": {
        "A": "Constituency Area",
        "B": "Labour Force Status",
        "F": "Population 15-24",
        "I": "NEET 15 -24",
        "L": "Population  18-35",
        "O": "NEET 18-35",
        "R": "Unemployment Rate",
        "S": "Not in Education, Employment and Training (NEET)",
    },
    "education": {
        "A": "Constituency Area",
        "B": "Literacy - Mother Tongue",
        "E": "Literacy - English",
        "G": "MDV Pop 15+",
        "H": "Highest Educational Attainment - Maldivians",
        "Q": "Highest Educational Attainment - Maldivians",
    },
}
DETAIL_HEADERS = {
    "population": {
        **dict(zip("BCDEFGHIJ", ("Total", "Male", "Female") * 3, strict=True)),
        **dict(
            zip(
                "MNOP",
                (
                    "Child Population (0-17) ",
                    "Local Youth Population ( 18-35 )",
                    "Working Age Population (36-64)",
                    "Elderly Population (65+)",
                ),
                strict=True,
            )
        ),
        **dict(
            zip(
                "QRST",
                (
                    "Child Population (0-17) ",
                    "Local Youth Population ( 18-35 )",
                    "Working Age Population (36-64)",
                    "Elderly Population (65+)",
                ),
                strict=True,
            )
        ),
    },
    "employment": {
        **dict(
            zip("BCDE", ("Employed", "Unemployed", "Outside Labour Force", "Total"), strict=True)
        ),
        **dict(zip("FGHIJKLMNOPQ", ("Male", "Female", "Total") * 4, strict=True)),
        "S": "15-24 years",
        "T": "18-35 years",
    },
    "education": {
        "B": "Resident Maldivian (10+)",
        "C": "Resident Maldivian  (Literate)",
        "D": "Literacy Rate (%)",
        "E": "Resident Maldivian  (Literate)",
        "F": "Literacy Rate (%)",
        **{get_column_letter(column): label for column, (_, label) in enumerate(ATTAINMENTS, 8)},
        **{
            get_column_letter(column): label
            for column, (_, label) in enumerate(ATTAINMENTS[1:], 17)
        },
    },
}


def extract_constituency_context(path: Path, kind: str) -> list[dict[str, Any]]:
    """Read every published metric cell without geographic/name-only deduplication."""
    if kind not in KINDS:
        raise ValueError(f"Unknown constituency table: {kind}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = kind.title()
        if workbook.sheetnames != [sheet_name]:
            raise ValueError(f"Constituency {kind} source sheet changed")
        sheet = workbook[sheet_name]
        values = list(sheet.values)
        metrics = METRICS[kind]
        if sheet.max_row != LAST_ROW or sheet.max_column != len(metrics) + 1:
            raise ValueError(f"Constituency {kind} row/column extent changed")
        if values[0][0] != TITLES[kind]:
            raise ValueError(f"Constituency {kind} source title changed")
        for source_row, expected in ((3, GROUP_HEADERS[kind]), (4, DETAIL_HEADERS[kind])):
            actual = {
                get_column_letter(column): value
                for column, value in enumerate(values[source_row - 1], 1)
                if value is not None
            }
            if actual != expected:
                raise ValueError(f"Constituency {kind} headers changed in row {source_row}")
        names = [row[0] for row in values[FIRST_ROW - 1 : LAST_ROW]]
        if any(not isinstance(name, str) or not name.strip() for name in names):
            raise ValueError(f"Constituency {kind} has an empty/non-text source name")
        repeated = Counter(names)
        source_headers = {}
        group = ""
        for column in range(2, len(metrics) + 2):
            letter = get_column_letter(column)
            group = GROUP_HEADERS[kind].get(letter, group)
            source_headers[letter] = " / ".join(
                label for label in (group, DETAIL_HEADERS[kind].get(letter)) if label
            )
        rows = []
        for source_row, values_row in enumerate(values[FIRST_ROW - 1 : LAST_ROW], FIRST_ROW):
            name = values_row[0]
            # These keys identify aligned rows in the three pinned workbooks, not places.
            key = f"census2022:constituency-row:{source_row}"
            for column, metric in enumerate(metrics, 2):
                if get_column_letter(column) != metric.column:
                    raise ValueError(f"Internal constituency {kind} metric order mismatch")
                published = values_row[column - 1]
                numeric = (
                    isinstance(published, (int, float))
                    and not isinstance(published, bool)
                    and math.isfinite(published)
                )
                if published is not None and not numeric:
                    raise ValueError(
                        f"Invalid source value at {sheet_name}!{metric.column}{source_row}"
                    )
                if numeric and (
                    published < 0
                    or (metric.published_unit == "ratio" and published > 1)
                    or (metric.published_unit == "percent" and published > 100)
                    or (metric.unit == "people" and int(published) != published)
                ):
                    raise ValueError(f"Out-of-range source value at {metric.column}{source_row}")
                value = (
                    None
                    if published is None
                    else (published * 100 if metric.published_unit == "ratio" else published)
                )
                rows.append(
                    {
                        "id": f"{key}:{kind}:{metric.metric}",
                        "constituency_key": key,
                        "constituency": name,
                        "geography_grain": "constituency",
                        "boundaryMatched": False,
                        "duplicate_source_name": repeated[name] > 1,
                        "year": YEAR,
                        "population_scope": metric.population_scope,
                        "sex": metric.sex,
                        "age_group": metric.age_group,
                        "metric": metric.metric,
                        "value": value,
                        "unit": metric.unit,
                        "denominator": metric.denominator,
                        "value_status": "reported" if numeric else "not_reported",
                        "published_value": published,
                        "published_unit": metric.published_unit,
                        "value_conversion": (
                            "ratio_to_percent_x100" if metric.published_unit == "ratio" else "none"
                        ),
                        "source_id": f"mbs-constituency-{kind}-2022",
                        "source_sheet": sheet_name,
                        "source_row": source_row,
                        "source_cell": f"{metric.column}{source_row}",
                        "source_header": source_headers[metric.column],
                    }
                )
        return rows
    finally:
        workbook.close()


def build_constituency_context(root: Path) -> dict[str, Any]:
    """Normalize pinned originals with checked receipts and no network dependency."""
    root = root.resolve()
    bundle = json.loads((root / "data/public/policy-evidence/bundle.json").read_text("utf-8"))
    sources_by_id = {source["id"]: source for source in bundle["sources"]}
    tables, sources, receipts = {}, [], []
    for kind in KINDS:
        source = sources_by_id[f"mbs-constituency-{kind}-2022"]
        relative = f"data/public/archive/policy-evidence/constituency-{kind}.xlsx"
        path = root / relative
        content = path.read_bytes()
        checksum = hashlib.sha256(content).hexdigest()
        if checksum != source["checksum"]:
            raise ValueError(f"Constituency {kind} source checksum differs from source register")
        receipt = {"path": relative, "sha256": checksum, "bytes": len(content)}
        tables[kind] = extract_constituency_context(path, kind)
        receipts.append(receipt)
        sources.append(
            {
                **source,
                "receipt": receipt,
                "modeledCoverage": "All 93 published constituency rows, including repeated names; "
                "2022 context with unmatched boundaries",
            }
        )
    # Positional keys are safe only after all pinned tables have the same row/name alignment.
    identities = [
        {(row["source_row"], row["constituency"]) for row in tables[kind]} for kind in KINDS
    ]
    if any(identity != identities[0] for identity in identities[1:]):
        raise ValueError("Constituency source rows differ between workbooks; review row keys")
    fingerprint = hashlib.sha256(
        (CONTEXT_VERSION + "".join(receipt["sha256"] for receipt in receipts)).encode()
        + Path(__file__).read_bytes()
    ).hexdigest()[:16]
    all_rows = [row for table in tables.values() for row in table]
    return {
        "schemaVersion": 1,
        "transformVersion": CONTEXT_VERSION,
        "version": f"{CONTEXT_VERSION}:{fingerprint}",
        "generatedAt": datetime.now(UTC).isoformat(),
        "sources": sources,
        "sourceReceipts": receipts,
        **tables,
        "coverage": {
            "constituency_source_rows": len(identities[0]),
            "distinct_source_names": len({name for _, name in identities[0]}),
            "first_source_row": FIRST_ROW,
            "last_source_row": LAST_ROW,
            "metric_cells": len(all_rows),
            "reported_values": sum(row["value_status"] == "reported" for row in all_rows),
            "blank_cells": sum(row["value_status"] == "not_reported" for row in all_rows),
            "ratio_cells_converted_to_percent": sum(
                row["value_conversion"] == "ratio_to_percent_x100" for row in all_rows
            ),
            "boundary_matched_rows": 0,
            "tables": {
                kind: {"rows": len(table), "metrics": len(METRICS[kind])}
                for kind, table in tables.items()
            },
        },
        "warnings": [
            GRAIN_WARNING,
            "There are 93 published rows but 92 distinct names: Meedhoo Dhaairaa appears at "
            "rows 69 and 97 with different values. Both rows are preserved, and names alone "
            "must not be used as join keys or grouping identities.",
            "Source ratios stored from 0 to 1 are multiplied by 100 for percent values. "
            "Population foreign share is already percent and is not multiplied again. "
            "Sex ratio is Maldivian males per 100 females, not a percent share.",
            "Education has blank count cells. Blanks remain null even when other published "
            "counts sum to the total; published ratios are retained rather than reconstructed.",
            "Population scopes, sex totals and age cohorts overlap. Filter to one metric "
            "before aggregation. Rates and sex ratios cannot be summed; means are unweighted.",
            "The 17 Greater Male population rows also exist in the policy-evidence bundle. "
            "These views share a source and must not be concatenated as independent records.",
        ],
    }
