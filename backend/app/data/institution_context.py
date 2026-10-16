"""Portable institutional history at the official source's annual metric grain.

Hospital activity is not unique patients, available appointments or resident demand.
Published stay/occupancy measures are retained, never recomputed from other rows.
"""

from __future__ import annotations

import hashlib
import json
import math
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

CONTEXT_VERSION = "official-institution-context-v1"
IGMH_SOURCE_ID = "mbs-igmh-utilisation-2025"
IGMH_RAW_PATH = "data/public/archive/policy-evidence/igmh-utilisation.xlsx"
IGMH_METRICS = (
    (
        "hospital_beds",
        "Number of hospital beds*",
        "beds at year-end",
        "Year-end beds; not average annual beds or immediately available beds. "
        "Do not add a second Dharumavantha campus total to this institutional series.",
    ),
    (
        "outpatient_activity",
        "out-patients",
        "reported encounters/year",
        "Published outpatient activity; repeat encounters may occur. "
        "Not unique patients or geographically assigned resident demand.",
    ),
    (
        "inpatient_activity",
        "In-patients",
        "reported admissions/year",
        "Published inpatient activity; not service-specific appointment availability.",
    ),
    (
        "inpatient_days",
        "In-patient Days",
        "patient-days/year",
        "Published institutional patient-days, retaining source precision.",
    ),
    (
        "average_duration_of_stay",
        "Average Duration of Stay (Days)",
        "days",
        "Reported mean stay; source cohorts/definitions can differ from patient-days "
        "divided by admissions, so the published value is not replaced.",
    ),
    (
        "average_occupancy_ratio",
        "Average Occupancy Ratio (%)",
        "percent",
        "Reported annual occupancy ratio; missing values are not calculated from "
        "year-end beds, admissions or patient-days.",
    ),
)


def extract_igmh_history(path: Path) -> list[dict[str, Any]]:
    """Read all cells, including explicit missing entries, with exact coordinates."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "6.6" not in workbook.sheetnames:
            raise ValueError("IGMH source sheet changed: expected 6.6")
        sheet = workbook["6.6"]
        values = list(sheet.values)
        if values[1][0] != (
            "Table 6.6 : IN - PATIENT HEALTH CARE IN INDIRA GANDHI MEMORIAL HOSPITAL,  2010 - 2025"
        ) or list(values[3][1:22]) != list(range(2005, 2026)):
            raise ValueError("IGMH title/year headers changed; review the source schema")
        if values[11][0] != (
            "Note:  *number of hospital beds are given according to the number on end of the year."
        ):
            raise ValueError("IGMH bed definition changed; review the source note")
        rows = []
        for source_row, (metric, label, unit, definition) in enumerate(IGMH_METRICS, 6):
            if values[source_row - 1][0] != label:
                raise ValueError(f"IGMH metric header changed at A{source_row}")
            for column, year in enumerate(range(2005, 2026), 2):
                published = values[source_row - 1][column - 1]
                numeric = (
                    isinstance(published, (int, float))
                    and not isinstance(published, bool)
                    and math.isfinite(published)
                )
                if numeric:
                    status = "reported"
                    if published < 0 or (metric == "average_occupancy_ratio" and published > 100):
                        raise ValueError(f"Invalid published {metric} for {year}: {published}")
                elif published is None:
                    status = "not_reported"
                elif isinstance(published, str) and published.strip().casefold() == "na":
                    status = "unavailable"
                else:
                    raise ValueError(f"Unrecognized IGMH source value: {published!r}")
                rows.append(
                    {
                        "id": f"igmh:{year}:{metric}",
                        "institution": "Indira Gandhi Memorial Hospital",
                        "geography_grain": "institution",
                        "year": year,
                        "metric": metric,
                        "value": published if numeric else None,
                        "unit": unit,
                        "value_status": status,
                        "published_value": published,
                        "source_id": IGMH_SOURCE_ID,
                        "source_table": "Table 6.6",
                        "source_sheet": "6.6",
                        "source_cell": f"{get_column_letter(column)}{source_row}",
                        "source_metric": label,
                        "definition": definition,
                        "qualification": "Historical institutional context; no current service "
                        "capacity, unique patient count, catchment or neighborhood allocation.",
                    }
                )
        return rows
    finally:
        workbook.close()


def build_institution_context(root: Path) -> dict[str, Any]:
    """Build from pinned local originals; never downloads or changes source files."""
    root = root.resolve()
    bundle_path = root / "data/public/policy-evidence/bundle.json"
    bundle = json.loads(bundle_path.read_text(encoding="utf-8"))
    source = next(item for item in bundle["sources"] if item["id"] == IGMH_SOURCE_ID)
    path = root / IGMH_RAW_PATH
    content = path.read_bytes()
    checksum = hashlib.sha256(content).hexdigest()
    if checksum != source["checksum"]:
        raise ValueError("IGMH source checksum differs from the captured source register")
    receipt = {"path": IGMH_RAW_PATH, "sha256": checksum, "bytes": len(content)}
    rows = extract_igmh_history(path)
    source_record = {
        **source,
        "modeledCoverage": "2005–2025 where reported; beds/outpatients begin in 2010",
        "receipt": receipt,
    }
    fingerprint = hashlib.sha256(
        (checksum + CONTEXT_VERSION).encode() + Path(__file__).read_bytes()
    ).hexdigest()[:16]
    return {
        "schemaVersion": 1,
        "transformVersion": CONTEXT_VERSION,
        "version": f"{CONTEXT_VERSION}:{fingerprint}",
        "generatedAt": datetime.now(UTC).isoformat(),
        "sources": [source_record],
        "sourceReceipts": [receipt],
        "igmh": rows,
        "coverage": {
            "institution_count": 1,
            "first_year": 2005,
            "last_year": 2025,
            "metrics": len(IGMH_METRICS),
            "cells": len(rows),
            "reported_values": sum(row["value_status"] == "reported" for row in rows),
            "blank_cells": sum(row["value_status"] == "not_reported" for row in rows),
            "explicitly_unavailable": sum(row["value_status"] == "unavailable" for row in rows),
        },
        "warnings": [
            "The published table title says 2010–2025; inpatient series also contain "
            "2005–2009 columns. Both the title and actual column years are preserved.",
            "The same 2024/2025 values also occur in policy observations; these views "
            "share a source and must not be concatenated as independent measurements.",
            "Reported mean stay does not always equal patient-days/admissions; retain "
            "reported measures until the publisher clarifies definitions.",
            "Occupancy in 2023/2024 is published as na; missing is not zero.",
            "This institutional series cannot establish available appointments or assign "
            "utilization to residents, origins or a separate hospital campus.",
        ],
    }
