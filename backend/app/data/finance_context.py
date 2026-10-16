"""Official programme budgets and MTCC aggregate accounts, never route cost factors."""

from __future__ import annotations

import hashlib
import json
import math
import re
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pypdf import PdfReader

VERSION = "official-finance-context-v2"
YEARS = (2024, 2025, 2026, 2027, 2028)
FISCAL_STATUS = {2024: "actual", 2025: "revised", 2026: "budget", 2027: "budget", 2028: "budget"}
SEGMENTS = (
    "trading",
    "construction_and_dredging",
    "transport",
    "real_estate",
    "others",
    "group_total",
)
FINANCE_ROWS = (
    (200, "total_revenue_before_intersegment", "Total revenue", "annual_flow"),
    (200, "intersegment_revenue_elimination", "Less: Inter-segment revenue", "annual_flow"),
    (200, "external_revenue", "Revenue", "annual_flow"),
    (200, "other_income", "Other income", "annual_flow"),
    (200, "operating_profit_loss", "Operating profit / (loss)", "annual_flow"),
    (200, "net_finance_cost", "Finance costs - net (Note 21)", "annual_flow"),
    (200, "profit_loss_before_tax", "Profit / (loss) before tax", "annual_flow"),
    (200, "income_tax", "Income tax", "annual_flow"),
    (200, "profit_loss_after_tax", "Profit / (loss) after tax", "annual_flow"),
    (201, "depreciation", "Depreciation (Note 6)", "annual_flow"),
    (
        201,
        "impairment_loss_financial_contract_assets",
        "Impairment loss on financial and contract assets (Note 11)",
        "annual_flow",
    ),
    (
        201,
        "obsolete_inventory_provision",
        "Provision for obsolete inventory (Note 19)",
        "annual_flow",
    ),
    (201, "amortisation", "Amortisation (Note 7)", "annual_flow"),
    (201, "assets", "Assets", "year_end_stock"),
    (201, "liabilities", "Liabilities", "year_end_stock"),
    (201, "capital_expenditure", "Capital expenditure (Note 6 and 7)", "annual_flow"),
)


def _amount(value: str) -> int | float:
    number = Decimal(value.replace(",", ""))
    if not number.is_finite():
        raise ValueError("Non-finite published finance amount")
    return int(number) if number == number.to_integral_value() else float(number)


def extract_programme_budget(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Retain the entire hierarchy, expose only its nonoverlapping terminal lines."""
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    table = soup.select_one("table#example")
    if table is None:
        raise ValueError("Budget programme table #example is absent")
    headers = [
        cell.get_text(" ", strip=True) for cell in table.select("tr")[0].find_all(["th", "td"])
    ]
    if headers != ["", "2024 Actual", "2025 Revised", "2026 Budget", "2027 Budget", "2028 Budget"]:
        raise ValueError("Budget fiscal year/status headers changed")
    nodes, stack = [], []
    for row_number, row in enumerate(table.select("tr[data-update-chart]"), 1):
        payload = json.loads(row["data-update-chart"])
        level = payload["level"]
        cells = row.find_all("td", recursive=False)
        if level != len(stack) and level > len(stack):
            raise ValueError("Budget hierarchy skips a parent level")
        stack = stack[:level]
        if len(cells) != 6 or row.get("data-level") != f"tr-level-{level}":
            raise ValueError("Budget row shape changed")
        amounts = {int(item["year"]): _amount(item["amount"]) for item in payload["years"]}
        if tuple(amounts) != YEARS:
            raise ValueError("Budget amount year columns changed")
        cell_ids = {}
        for year, cell in zip(YEARS, cells[1:], strict=True):
            if _amount(cell.get_text("", strip=True)) != round(amounts[year]):
                raise ValueError("Budget displayed amount disagrees with exact chart payload")
            cell_ids[str(year)] = cell["id"]
        node_id = f"budget:{cells[1]['id'].removeprefix('td-2024-')}"
        node = {
            "id": node_id,
            "programme": payload["name"],
            "level": level,
            "parent_id": stack[-1]["id"] if stack else None,
            "path": [item["programme"] for item in stack] + [payload["name"]],
            "source_table": "Organizations - Special Budget / table#example",
            "source_row": row_number,
            "source_cell_ids": cell_ids,
            "amounts": {str(year): amount for year, amount in amounts.items()},
        }
        nodes.append(node)
        stack.append(node)
    if not nodes or nodes[0]["programme"] != "Special Budget" or nodes[0]["level"] != 0:
        raise ValueError("Expected Special Budget root; this is not the total national budget")
    parent_ids = {node["parent_id"] for node in nodes if node["parent_id"]}
    allocations = []
    children: dict[str, list[dict[str, Any]]] = {}
    for node in nodes:
        children.setdefault(node["parent_id"], []).append(node)
        node["is_leaf"] = node["id"] not in parent_ids
        if not node["is_leaf"]:
            continue
        for year in YEARS:
            allocations.append(
                {
                    "id": f"{node['id']}:{year}",
                    "node_id": node["id"],
                    "programme": node["programme"],
                    "programme_path": " / ".join(node["path"]),
                    "programme_group": node["path"][1] if len(node["path"]) > 1 else None,
                    "year": year,
                    "fiscal_status": FISCAL_STATUS[year],
                    "value": node["amounts"][str(year)],
                    "unit": "MVR",
                    "currency": "MVR",
                    "geography_grain": "government_programme_line",
                    "source_id": "budget2026",
                    "source_table": node["source_table"],
                    "source_row": node["source_row"],
                    "source_cell": node["source_cell_ids"][str(year)],
                    "qualification": "Captured Special Budget terminal line; actual/revised/budget "
                    "statuses differ. No route, island, per-resident or marginal cost allocation.",
                }
            )
    # A terminal-line view must reconcile to each published parent, not duplicate its totals.
    for node in nodes:
        if node["is_leaf"]:
            continue
        for year in YEARS:
            subtotal = sum(
                Decimal(str(child["amounts"][str(year)])) for child in children[node["id"]]
            )
            if subtotal != Decimal(str(node["amounts"][str(year)])):
                raise ValueError(f"Budget children do not reconcile: {node['id']} {year}")
    return nodes, allocations


def _financial_tokens(text: str) -> list[str]:
    return re.findall(r"\([\d,]+\)|\b\d[\d,]*\b|(?<!\S)-(?!\S)", text)


def extract_mtcc_context(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    reader = PdfReader(path)
    if len(reader.pages) != 206:
        raise ValueError("MTCC annual report pagination changed; review table positions")
    texts = {page: reader.pages[page - 1].extract_text() or "" for page in (65, 66, 67, 200, 201)}
    finances = []
    for page, metric, source_label, measure_type in FINANCE_ROWS:
        text = texts[page]
        if "29 Segment information - Group" not in text or "Maldivian Rufiyaa" not in text:
            raise ValueError("Expected audited MTCC Group Note 29 in MVR")
        matches = []
        for line in text.splitlines():
            normal = " ".join(line.split())
            if normal.startswith(source_label + " "):
                matches.append(_financial_tokens(normal[len(source_label) :]))
        if len(matches) != 2 or any(len(tokens) != 6 for tokens in matches):
            raise ValueError(f"MTCC segment row/column schema changed: {source_label}")
        for year, tokens in zip((2025, 2024), matches, strict=True):
            for segment, published in zip(SEGMENTS, tokens, strict=True):
                value = 0 if published == "-" else _amount(published.strip("()"))
                if published.startswith("("):
                    value = -value
                finances.append(
                    {
                        "id": f"mtcc:group:{year}:{segment}:{metric}",
                        "entity": "MTCC Group",
                        "segment": segment,
                        "year": year,
                        "metric": metric,
                        "value": value,
                        "unit": "MVR",
                        "currency": "MVR",
                        "measure_type": measure_type,
                        "geography_grain": "company_group_business_segment",
                        "is_group_total": segment == "group_total",
                        "published_value": published,
                        "value_status": "reported_nil" if published == "-" else "reported",
                        "source_id": "mtcc-annual2025",
                        "source_page": page,
                        "source_statement_page": page - 125,
                        "source_table": "Audited financial statements Note 29: "
                        "Segment information - Group",
                        "source_row": source_label,
                        "source_column": segment,
                        "qualification": "Group segment account, not route operating or "
                        "marginal cost. Group total overlaps its components; metrics and "
                        "stock/flow measures differ. "
                        "Published accounting dashes are retained as reported nil (zero).",
                    }
                )
    for year in (2024, 2025):
        for _, metric, _, _ in FINANCE_ROWS:
            rows = [row for row in finances if row["year"] == year and row["metric"] == metric]
            total = next(row["value"] for row in rows if row["is_group_total"])
            components = sum(row["value"] for row in rows if not row["is_group_total"])
            if abs(total - components) > 1:
                raise ValueError(f"MTCC segment amounts do not reconcile: {metric} {year}")
    passengers = []
    bus, ferry = texts[65].split("RTL Ferry Ridership", 1)
    bus = bus.split("RTL Bus Ridership", 1)[0]
    series = (
        ("rtl_bus_male_region", "Greater Malé RTL buses", 65, bus),
        ("rtl_ferry_all_zones", "RTL ferries, all published zones", 65, ferry),
        ("all_modes", "MTCC transport division, all modes", 67, texts[67]),
    )
    for series_id, scope, page, text in series:
        values = re.findall(r"\b\d{1,3}(?:,\d{3})+\b", text)
        if len(values) != 4:
            raise ValueError(f"MTCC annual passenger chart changed: {series_id}")
        for year, published in zip(range(2022, 2026), values, strict=True):
            passengers.append(_passenger_row(series_id, scope, year, published, page))
    zones = re.findall(r"RTL Zone (\d)\s*-\s*RTLZ\d\s+([\d,]+)", texts[66])
    if {zone for zone, _ in zones} != {"1", "2", "4", "5", "6"} or len(zones) != 5:
        raise ValueError("MTCC ferry zone chart changed")
    for zone, published in zones:
        passengers.append(
            _passenger_row(f"rtl_ferry_zone_{zone}", f"RTL ferry Zone {zone}", 2025, published, 66)
        )
    ferry_total = next(
        row["value"]
        for row in passengers
        if row["series"] == "rtl_ferry_all_zones" and row["year"] == 2025
    )
    if sum(row["value"] for row in passengers if row["source_page"] == 66) != ferry_total:
        raise ValueError("MTCC ferry zone totals do not reconcile")
    return finances, passengers


def _passenger_row(series: str, scope: str, year: int, published: str, page: int) -> dict[str, Any]:
    return {
        "id": f"mtcc:passengers:{series}:{year}",
        "series": series,
        "scope": scope,
        "year": year,
        "metric": "reported_passengers_carried",
        "value": _amount(published),
        "unit": "reported passenger movements/year",
        "published_value": published,
        "geography_grain": "operator_service_category",
        "source_id": "mtcc-annual2025",
        "source_page": page,
        "source_table": "Transport Services Division annual ridership charts",
        "qualification": "Annual operator-reported movements, not unique people, route boardings "
        "or verified recurring service. All modes overlaps bus/ferry; ferry totals overlap zones. "
        "MBS monthly publications may report overlapping categories with differing definitions.",
    }


def extract_hies_expenditure(path: Path) -> list[dict[str, Any]]:
    """Keep every published estimate, without treating means/shares as additive totals."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["expenditure"]:
            raise ValueError("HIES expenditure worksheet changed")
        sheet = workbook["expenditure"]
        rows = list(sheet.values)
        if (
            len(rows) != 74
            or rows[0][0] != "Household Income and Expenditure Surveys 2019"
            or rows[2][1] != 2019
            or rows[3][1] != "Republic"
            or not str(rows[3][2]).startswith("Male")
            or rows[3][3] != "Atolls"
            or rows[4][0] != "Expenditure (for expenditure reported households)"
        ):
            raise ValueError("HIES year, geography or table schema changed")
        sections = (
            (6, 7, 13, "summary", "Average household and per capita expenditure ( in MVR)"),
            (14, 15, 28, "category_share", "Expenditure share by broad expenditure groups (in %)"),
            (
                29,
                30,
                42,
                "household_category",
                "Average monthly household expenditure by broad expenditure groups (in MVR)",
            ),
            (
                43,
                44,
                56,
                "per_capita_category",
                "Average monthly per capita expenditure by broad expenditure groups (in MVR)",
            ),
            (
                57,
                58,
                68,
                "per_capita_by_household_size",
                "Average monthly per capita expenditure by household size(in MVR)",
            ),
            (
                69,
                70,
                74,
                "per_capita_by_quintile",
                "Per capita monthly expenditure by per capita quintiles (in MVR)",
            ),
        )
        expected_summary = (
            "Average Household Size",
            "Average Monthly Expenditure per Household",
            "Median Monthly Expenditure per Household",
            "Monthly per capita Expenditure",
            "Median monthly per capita expenditure",
            "Expenditure per person per day",
            "Median of Expenditure per person per day",
        )
        if tuple(str(rows[row - 1][0]).strip() for row in range(7, 14)) != expected_summary:
            raise ValueError("HIES summary indicator labels changed")
        # Three category presentations must use the same ordered 13 broad groups.
        category_labels = [str(rows[row - 1][0]).strip() for row in range(30, 43)]
        share_labels = [str(rows[row - 1][0]).strip() for row in range(15, 29) if row != 19]
        per_capita_labels = [str(rows[row - 1][0]).strip() for row in range(44, 57)]
        if category_labels != share_labels or category_labels != per_capita_labels:
            raise ValueError("HIES broad expenditure category columns changed")
        output = []
        for header, first, last, section, heading in sections:
            if rows[header - 1][0] != heading:
                raise ValueError(f"HIES expenditure section changed: {section}")
            for source_row in range(first, last + 1):
                source_indicator = rows[source_row - 1][0]
                indicator = str(source_indicator).strip()
                if section == "category_share":
                    unit = "% of household expenditure"
                elif source_row == 7:
                    unit = "persons/household"
                elif source_row in (8, 9) or section == "household_category":
                    unit = "MVR/household/month"
                elif source_row in (12, 13):
                    unit = "MVR/person/day"
                else:
                    unit = "MVR/person/month"
                for column, geography in enumerate(("Republic", "Malé", "Atolls"), 2):
                    published = rows[source_row - 1][column - 1]
                    if (
                        isinstance(published, bool)
                        or not isinstance(published, (int, float))
                        or not math.isfinite(published)
                        or published < 0
                    ):
                        raise ValueError(
                            f"HIES estimate is not finite/nonnegative: row {source_row}"
                        )
                    if section == "category_share" and not 0 <= published <= 1:
                        raise ValueError("HIES published share is not a fraction")
                    output.append(
                        {
                            "id": f"hies:2019:{column}:{source_row}",
                            "year": 2019,
                            "geography": geography,
                            "published_geography": rows[3][column - 1],
                            "geography_grain": "national_or_capital_or_atoll_aggregate",
                            "section": section,
                            "section_label": heading,
                            "indicator": indicator,
                            "source_indicator": source_indicator,
                            "value": published * 100 if section == "category_share" else published,
                            "published_value": published,
                            "value_conversion": "fraction_times_100"
                            if section == "category_share"
                            else "identity",
                            "unit": unit,
                            "currency": None
                            if source_row == 7 or section == "category_share"
                            else "MVR",
                            "price_basis": "nominal 2019" if unit.startswith("MVR") else None,
                            "measure_type": "published_survey_estimate_nonadditive",
                            "is_subcategory": source_row == 19,
                            "source_id": "mbs-hies-expenditure-2019",
                            "source_sheet": sheet.title,
                            "source_row": source_row,
                            "source_cell": f"{chr(64 + column)}{source_row}",
                            "qualification": "HIES 2019 reported-expenditure household "
                            "survey estimate. Means, medians and shares are not additive. "
                            "Republic overlaps Malé/Atolls; no island allocation or "
                            "current-price interpretation. Housing rent is a subset of "
                            "housing expenditure, not an additional broad category. "
                            "Category means must not be summed to overall expenditure "
                            "without source weights.",
                        }
                    )
        return output
    finally:
        workbook.close()


def build_finance_context(root: Path) -> dict[str, Any]:
    root = root.resolve()
    raw = root / "data/raw"
    manifest = json.loads((raw / "source-manifest.json").read_text(encoding="utf-8"))
    indexed = {item["id"]: item for item in manifest["sources"]}
    sources, receipts = [], []
    for identifier in ("budget2026", "mtcc-annual2025"):
        source = indexed[identifier]
        path = raw / source["file"]
        content = path.read_bytes()
        receipt = {
            "path": path.relative_to(root).as_posix(),
            "sha256": hashlib.sha256(content).hexdigest(),
            "bytes": len(content),
        }
        if receipt["sha256"] != source["sha256"]:
            raise ValueError(
                f"Finance source checksum differs from captured manifest: {identifier}"
            )
        receipts.append(receipt)
        sources.append({**source, "receipt": receipt})
    nodes, allocations = extract_programme_budget(raw / indexed["budget2026"]["file"])
    finances, passengers = extract_mtcc_context(raw / indexed["mtcc-annual2025"]["file"])
    policy = json.loads(
        (root / "data/public/policy-evidence/bundle.json").read_text(encoding="utf-8")
    )
    for source in policy["sources"]:
        if source["id"] not in ("mbs-hies-expenditure-2019", "mbs-hies-expenditure-methods-2019"):
            continue
        path = root / source["rawPath"]
        content = path.read_bytes()
        receipt = {
            "path": source["rawPath"],
            "sha256": hashlib.sha256(content).hexdigest(),
            "bytes": len(content),
        }
        if receipt["sha256"] != source["checksum"]:
            raise ValueError(f"HIES source checksum differs from captured manifest: {source['id']}")
        sources.append({**source, "receipt": receipt})
        receipts.append(receipt)
    if len(sources) != 4:
        raise ValueError("Both HIES workbook and methods source receipts are required")
    hies_source = next(source for source in sources if source["id"] == "mbs-hies-expenditure-2019")
    hies = extract_hies_expenditure(root / hies_source["rawPath"])
    fingerprint = hashlib.sha256(
        json.dumps(receipts, sort_keys=True).encode() + Path(__file__).read_bytes()
    ).hexdigest()[:16]
    return {
        "schemaVersion": 1,
        "transformVersion": VERSION,
        "version": f"{VERSION}:{fingerprint}",
        "generatedAt": datetime.now(UTC).isoformat(),
        "sources": sources,
        "sourceReceipts": receipts,
        "budgetProgrammeHierarchy": nodes,
        "budgetAllocations": allocations,
        "mtccSegmentAccounts": finances,
        "mtccPassengerAnnual": passengers,
        "hiesExpenditure": hies,
        "coverage": {
            "budget_hierarchy_nodes": len(nodes),
            "budget_terminal_lines": len(allocations) // len(YEARS),
            "budget_allocation_year_rows": len(allocations),
            "budget_years": list(YEARS),
            "mtcc_segment_account_rows": len(finances),
            "mtcc_passenger_rows": len(passengers),
            "mtcc_financial_years": [2024, 2025],
            "hies_expenditure_rows": len(hies),
            "hies_expenditure_indicators": len(hies) // 3,
            "hies_reference_year": 2019,
        },
        "warnings": [
            "Budget capture is Organizations - Special Budget, not the whole national budget. "
            "The terminal lines reconcile to its published parent totals.",
            "2024 is actual, 2025 revised, and 2026–2028 budget; appropriations do not prove "
            "spending, delivery, service frequency or current route subsidy receipts.",
            "Programme hierarchy totals overlap their descendants; only terminal lines are "
            "exposed as additive programme allocations.",
            "MTCC figures are Group Note 29 accounts; Company tables and rounded management "
            "charts are separate presentations, not additional independent measurements.",
            "Business-segment profits include accounting allocations and operational support; "
            "no aggregate account has been converted into cost per trip, route, passenger or km.",
            "MTCC annual passenger categories overlap one another and potentially MBS monthly "
            "series; source definitions must be reconciled before any combined series is made.",
            "HIES 2019 estimates cover Republic, Malé and Atolls aggregates only. Means, "
            "medians and shares are nonadditive; percentages explicitly convert published "
            "fractions. Nominal 2019 amounts are not current local travel affordability.",
        ],
    }
