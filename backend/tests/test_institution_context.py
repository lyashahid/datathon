"""Institution history preserves source grain, missingness and reproducibility."""

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.analytics.models import ChartRequest, DatasetFilter, DatasetQuery
from app.analytics.service import DatasetService
from app.data.institution_context import (
    IGMH_METRICS,
    build_institution_context,
    extract_igmh_history,
)

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(scope="module")
def bundle():
    return json.loads((ROOT / "data/public/institution-context.json").read_text(encoding="utf-8"))


def test_history_retains_every_published_cell_and_real_periods(bundle):
    rows = bundle["igmh"]
    assert len(rows) == len({row["id"] for row in rows}) == 126
    assert {row["year"] for row in rows} == set(range(2005, 2026))
    assert len({(row["metric"], row["year"]) for row in rows}) == 126
    assert bundle["coverage"]["reported_values"] == 114
    assert bundle["coverage"]["blank_cells"] == 10
    assert bundle["coverage"]["explicitly_unavailable"] == 2
    assert all(row["geography_grain"] == "institution" for row in rows)
    blank = [row for row in rows if row["value_status"] == "not_reported"]
    assert {row["metric"] for row in blank} == {"hospital_beds", "outpatient_activity"}
    assert {row["year"] for row in blank} == set(range(2005, 2010))
    assert all(row["value"] is None and row["published_value"] is None for row in blank)


def test_reported_precision_missing_occupancy_and_nonreconciling_stay_are_preserved(bundle):
    by_key = {(row["metric"], row["year"]): row for row in bundle["igmh"]}
    assert by_key["inpatient_days", 2020]["value"] == 66244.63
    assert by_key["hospital_beds", 2025]["value"] == 537
    assert by_key["outpatient_activity", 2025]["value"] == 464783
    assert by_key["hospital_beds", 2025]["source_cell"] == "V6"
    for year in (2023, 2024):
        occupancy = by_key["average_occupancy_ratio", year]
        assert occupancy["value"] is None
        assert occupancy["published_value"] == "na"
        assert occupancy["value_status"] == "unavailable"
    reported_stay = by_key["average_duration_of_stay", 2023]["value"]
    assert reported_stay == 4.7
    assert reported_stay != pytest.approx(
        by_key["inpatient_days", 2023]["value"] / by_key["inpatient_activity", 2023]["value"]
    )


def test_source_schema_drift_fails_instead_of_silently_shifting_metrics(tmp_path):
    path = tmp_path / "wrong.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "6.6"
    sheet.cell(2, 1, "changed title")
    for column, year in enumerate(range(2005, 2026), 2):
        sheet.cell(4, column, year)
    for row, (_, label, _, _) in enumerate(IGMH_METRICS, 6):
        sheet.cell(row, 1, label)
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="title/year headers changed"):
        extract_igmh_history(path)


def test_portable_history_is_queryable_without_original_workbook(bundle, tmp_path):
    public = tmp_path / "public"
    public.mkdir()
    (public / "institution-context.json").write_text(json.dumps(bundle), encoding="utf-8")
    service = DatasetService(public, tmp_path)
    schema = service.schema("igmh-utilisation-history")
    assert schema["row_count"] == 126
    assert schema["provenance"]["sources"][0]["receipt"]["sha256"]
    selected = service.query(
        "igmh-utilisation-history",
        DatasetQuery(
            filters=[
                DatasetFilter(column="metric", value="average_occupancy_ratio"),
                DatasetFilter(column="year", op="gte", value=2023),
            ],
            sort="year",
        ),
    )
    assert [row["value"] for row in selected["items"]] == [None, None, 45.9]
    chart = service.chart(
        "igmh-utilisation-history",
        ChartRequest(
            group_by="year",
            measure="value",
            aggregation="sum",
            filters=[DatasetFilter(column="metric", value="average_occupancy_ratio")],
        ),
    )
    assert next(row for row in chart["items"] if row["group"] == "2024")["value"] is None
    with pytest.raises(ValueError, match="cannot be summed across years"):
        service.chart(
            "igmh-utilisation-history",
            ChartRequest(group_by="metric", measure="value", aggregation="sum"),
        )
    with pytest.raises(ValueError, match="Mixed units"):
        service.chart(
            "igmh-utilisation-history",
            ChartRequest(group_by="year", measure="value", aggregation="mean"),
        )


@pytest.mark.skipif(
    not (ROOT / "data/public/archive/policy-evidence/igmh-utilisation.xlsx").is_file(),
    reason="Optional official workbook capture is not installed",
)
def test_captured_workbook_reproduces_every_artifact_value(bundle):
    rebuilt = build_institution_context(ROOT)
    for key in ("igmh", "coverage", "sources", "sourceReceipts", "warnings"):
        assert rebuilt[key] == bundle[key]
