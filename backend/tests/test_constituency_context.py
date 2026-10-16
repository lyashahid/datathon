"""National constituency cells stay traceable, correctly scaled and geographically isolated."""

import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.utils import column_index_from_string

from app.analytics.models import ChartRequest, DatasetFilter, DatasetQuery
from app.analytics.service import DatasetService
from app.data.constituency_context import (
    DETAIL_HEADERS,
    FIRST_ROW,
    GROUP_HEADERS,
    KINDS,
    LAST_ROW,
    METRICS,
    TITLES,
    build_constituency_context,
    extract_constituency_context,
)

ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = ROOT / "data/public/archive/policy-evidence"


@pytest.fixture(scope="module")
def bundle():
    return json.loads((ROOT / "data/public/constituency-context.json").read_text("utf-8"))


def _cell(bundle, kind, cell):
    return next(row for row in bundle[kind] if row["source_cell"] == cell)


def test_all_national_cells_and_last_constituency_are_retained(bundle):
    coverage = bundle["coverage"]
    assert coverage["constituency_source_rows"] == 93
    assert coverage["distinct_source_names"] == 92
    assert coverage["metric_cells"] == 5673
    assert coverage["reported_values"] == 5645
    assert coverage["blank_cells"] == 28
    assert coverage["ratio_cells_converted_to_percent"] == 1581
    ids = set()
    for kind in KINDS:
        rows = bundle[kind]
        assert len(rows) == 93 * len(METRICS[kind])
        assert {row["source_row"] for row in rows} == set(range(FIRST_ROW, LAST_ROW + 1))
        assert {row["source_cell"] for row in rows} == {
            f"{metric.column}{source_row}"
            for source_row in range(FIRST_ROW, LAST_ROW + 1)
            for metric in METRICS[kind]
        }
        assert all(row["year"] == 2022 for row in rows)
        assert all(row["source_sheet"] == kind.title() for row in rows)
        assert _cell(bundle, kind, "B97")["constituency"] == "Meedhoo Dhaairaa"
        ids.update(row["id"] for row in rows)
    assert len(ids) == coverage["metric_cells"]


def test_repeated_names_have_distinct_source_row_keys_and_no_island_join(bundle):
    rows = [row for row in bundle["population"] if row["metric"] == "resident_population"]
    meedhoo = [row for row in rows if row["constituency"] == "Meedhoo Dhaairaa"]
    assert [(row["source_row"], row["value"]) for row in meedhoo] == [(69, 2727), (97, 1880)]
    assert len({row["constituency_key"] for row in meedhoo}) == 2
    assert all(row["duplicate_source_name"] for row in meedhoo)
    for kind in KINDS:
        assert all(row["geography_grain"] == "constituency" for row in bundle[kind])
        assert all(row["boundaryMatched"] is False for row in bundle[kind])
        assert all("island_id" not in row and "islandId" not in row for row in bundle[kind])
    assert bundle["coverage"]["boundary_matched_rows"] == 0


def test_ratio_percent_and_sex_ratio_units_are_not_conflated(bundle):
    cases = (("population", "Q5"), ("employment", "R5"), ("education", "D5"))
    for kind, coordinate in cases:
        row = _cell(bundle, kind, coordinate)
        assert row["published_unit"] == "ratio"
        assert row["unit"] == "percent"
        assert row["value_conversion"] == "ratio_to_percent_x100"
        assert row["value"] == row["published_value"] * 100
    assert _cell(bundle, "employment", "R5")["value"] == pytest.approx(5.0952045133991535)
    foreign = _cell(bundle, "population", "K5")
    assert foreign["published_unit"] == foreign["unit"] == "percent"
    assert foreign["value"] == foreign["published_value"] == 22.11676340047798
    assert foreign["value_conversion"] == "none"
    sex_ratio = _cell(bundle, "population", "L5")
    assert sex_ratio["unit"] == "Maldivian males per 100 females"
    assert sex_ratio["value"] == pytest.approx(5629 / 5777 * 100)
    assert sex_ratio["value_conversion"] == "none"
    assert _cell(bundle, "population", "O5")["age_group"] == "36-64"
    assert _cell(bundle, "education", "D5")["age_group"] == "10+"
    assert _cell(bundle, "education", "Q5")["age_group"] == "15+"


def test_missing_education_counts_are_not_filled_from_published_zero_ratios(bundle):
    blank = _cell(bundle, "education", "I35")
    assert blank["metric"] == "attainment_preschool"
    assert blank["value"] is None and blank["published_value"] is None
    assert blank["value_status"] == "not_reported"
    ratio = _cell(bundle, "education", "Q35")
    assert ratio["value"] == ratio["published_value"] == 0
    assert ratio["value_status"] == "reported"
    assert _cell(bundle, "education", "H5")["metric"] == "attainment_not_stated"
    assert _cell(bundle, "education", "O5")["metric"] == "attainment_level_not_stated"
    assert sum(row["value"] is None for row in bundle["education"]) == 28


def _fixture_workbook(path, kind, changed_cell=None, changed_value=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = kind.title()
    sheet["A1"] = TITLES[kind]
    for source_row, headers in ((3, GROUP_HEADERS[kind]), (4, DETAIL_HEADERS[kind])):
        for column, label in headers.items():
            sheet[f"{column}{source_row}"] = label
    for source_row in range(FIRST_ROW, LAST_ROW + 1):
        sheet.cell(source_row, 1, f"Constituency {source_row}")
        for metric in METRICS[kind]:
            sheet.cell(source_row, column_index_from_string(metric.column), 0)
    if changed_cell:
        sheet[changed_cell] = changed_value
    workbook.save(path)
    workbook.close()


@pytest.mark.parametrize(
    ("kind", "cell", "value", "message"),
    (
        ("education", "D4", "Different measure", "headers changed"),
        ("population", "A1", "Census 2030", "title changed"),
        ("employment", "A97", None, "empty/non-text"),
        ("employment", "R5", 5.1, "Out-of-range"),
        ("population", "K5", 100.1, "Out-of-range"),
        ("education", "B5", 2.5, "Out-of-range"),
        ("employment", "B5", True, "Invalid source value"),
        ("education", "B5", "#DIV/0!", "Invalid source value"),
    ),
)
def test_schema_drift_and_invalid_units_fail_explicitly(tmp_path, kind, cell, value, message):
    path = tmp_path / f"{kind}.xlsx"
    _fixture_workbook(path, kind, cell, value)
    with pytest.raises(ValueError, match=message):
        extract_constituency_context(path, kind)


def test_changed_source_bytes_are_rejected(tmp_path):
    policy = tmp_path / "data/public/policy-evidence"
    policy.mkdir(parents=True)
    archive = tmp_path / "data/public/archive/policy-evidence"
    archive.mkdir(parents=True)
    (archive / "constituency-population.xlsx").write_bytes(b"changed capture")
    (policy / "bundle.json").write_text(
        json.dumps({"sources": [{"id": "mbs-constituency-population-2022", "checksum": "wrong"}]}),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="checksum differs"):
        build_constituency_context(tmp_path)


def test_portable_models_are_queryable_with_metric_and_rate_aggregation_guards(bundle, tmp_path):
    public = tmp_path / "public"
    public.mkdir()
    (public / "constituency-context.json").write_text(json.dumps(bundle), encoding="utf-8")
    service = DatasetService(public, tmp_path)
    for kind in KINDS:
        schema = service.schema(f"constituency-{kind}-context")
        assert schema["row_count"] == bundle["coverage"]["tables"][kind]["rows"]
        assert len(schema["provenance"]["sources"]) == 3
        assert all(source["receipt"]["sha256"] for source in schema["provenance"]["sources"])
    selected = service.query(
        "constituency-population-context",
        DatasetQuery(
            filters=[
                DatasetFilter(column="constituency", value="Meedhoo Dhaairaa"),
                DatasetFilter(column="metric", value="resident_population"),
            ],
            sort="source_row",
        ),
    )
    assert [row["value"] for row in selected["items"]] == [2727, 1880]
    dataset = "constituency-employment-context"
    rate_filter = [DatasetFilter(column="metric", value="unemployment_percent")]
    chart = service.chart(
        dataset,
        ChartRequest(
            group_by="constituency_key",
            measure="value",
            aggregation="mean",
            filters=rate_filter,
        ),
    )
    assert len(chart["items"]) == 93
    assert (
        next(row for row in chart["items"] if row["group"].endswith(":5"))["value"]
        == (_cell(bundle, "employment", "R5")["value"])
    )
    with pytest.raises(ValueError, match="rates and sex ratios cannot be summed"):
        service.chart(
            dataset,
            ChartRequest(
                group_by="constituency_key",
                measure="value",
                aggregation="sum",
                filters=rate_filter,
            ),
        )
    with pytest.raises(ValueError, match="metric filter or grouping"):
        service.chart(
            dataset,
            ChartRequest(
                group_by="constituency_key",
                measure="value",
                aggregation="mean",
            ),
        )
    with pytest.raises(ValueError, match="Repeated constituency names"):
        service.chart(
            dataset,
            ChartRequest(
                group_by="constituency",
                measure="value",
                aggregation="mean",
                filters=rate_filter,
            ),
        )
    count_chart = service.chart(
        dataset,
        ChartRequest(
            group_by="constituency_key",
            measure="value",
            aggregation="sum",
            filters=[DatasetFilter(column="metric", value="employed")],
        ),
    )
    assert len(count_chart["items"]) == 93


@pytest.mark.skipif(
    not all((RAW_DIR / f"constituency-{kind}.xlsx").is_file() for kind in KINDS),
    reason="Optional official workbook captures are not installed",
)
def test_original_workbooks_reproduce_all_values_and_receipts(bundle):
    rebuilt = build_constituency_context(ROOT)
    for key in (*KINDS, "coverage", "sources", "sourceReceipts", "warnings", "version"):
        assert rebuilt[key] == bundle[key]
